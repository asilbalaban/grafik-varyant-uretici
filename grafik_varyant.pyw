# -*- coding: utf-8 -*-
"""Grafik Varyant Üretici

Şablon dosyalarındaki (SVG; Illustrator ile .ai/.eps; Illustrator'süz PDF) isimlendirilmiş yazı
alanlarını (ör. NAME_S, NUMBER_M) bir tablodaki verilerle doldurup her satır için ayrı dosya üretir.
"""
import csv
import datetime
import html
import io
import json
import os
import re
import struct
import sys
import tempfile
from pathlib import Path

import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from tkinter import font as tkfont

APP = "Grafik Varyant Üretici"
COL_T = "Beden"  # her satırın hangi şablon dosyasıyla üretileceğini söyleyen sütun
COL_T_OLD = "Şablon"
COL_F = "Dosya Adı"
CHECK_DIR = "KONTROL-EDILECEK"  # elle bakılması gereken çıktıların ayrıldığı alt klasör
SIZE_ALIASES = [("2xl", "xxl"), ("3xl", "xxxl"), ("4xl", "xxxxl"), ("5xl", "xxxxxl"), ("2xs", "xxs"), ("3xs", "xxxs")]
SKIP = "(yok say)"

CASES = {"asis": "Olduğu gibi", "upper": "BÜYÜK HARF", "lower": "küçük harf", "title": "Baş Harfler Büyük"}
EMPTY = {"blank": "Boş bırak", "keep": "Şablondaki yazı kalsın"}
ALIGN = {"left": "Sola yaslı", "center": "Ortalı", "right": "Sağa yaslı"}
COLLISION = {"overwrite": "Üzerine yaz", "skip": "Atla (mevcut dosya kalsın)", "number": "Numara ekle (_2, _3…)"}

ALIAS_T = ["şablon", "beden", "size", "kalıp", "template", "varyant", "model"]
ALIAS_F = ["dosya adı", "dosya", "dosya ismi", "filename", "file", "file name"]
ALIAS_GROUPS = [
    ["name", "isim", "ad", "ad soyad", "adsoyad", "isim soyisim"],
    ["number", "numara", "no", "num", "nr", "sayı"],
    ["blood", "kan", "kan grubu", "blood type", "bloodtype"],
]

# ----------------------------------------------------------------------------
# Çekirdek: metin yardımcıları
# ----------------------------------------------------------------------------


def tr_lower(s):
    return s.replace("İ", "i").replace("I", "ı").lower()


def tr_upper(s):
    return s.replace("i", "İ").replace("ı", "I").upper()


def tr_title(s):
    return " ".join(tr_upper(w[:1]) + tr_lower(w[1:]) for w in s.split(" "))


def norm(s):
    """Karşılaştırma anahtarı: büyük/küçük harf, boşluk, _ ve - farklarını yok sayar."""
    return re.sub(r"[\s_\-]+", "", tr_lower(str(s).strip())).replace("ı", "i")


def fmt_num(x):
    return f"{x:.2f}".rstrip("0").rstrip(".")


def xml_escape(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ----------------------------------------------------------------------------
# Çekirdek: SVG okuma
# ----------------------------------------------------------------------------

TAG_RE = re.compile(
    r"<!--.*?-->|<!\[CDATA\[.*?\]\]>|<\?.*?\?>|"
    r"<(/?)([A-Za-z_][\w:.\-]*)((?:[^<>\"']|\"[^\"]*\"|'[^']*')*)>",
    re.S,
)
TSPAN_RE = re.compile(r"<tspan\b(?:[^<>\"']|\"[^\"]*\"|'[^']*')*>")
STYLE_ATTR_RE = re.compile(r"(?<![\w:\-])style\s*=\s*([\"'])(.*?)\1", re.S)
SIZE_RE = re.compile(r"^\s*([\d.]+)\s*([a-zA-Z%]*)")


def get_attr(attrs, name):
    m = re.search(r"(?<![\w:\-])%s\s*=\s*(?:\"([^\"]*)\"|'([^']*)')" % re.escape(name), attrs)
    if not m:
        return None
    return html.unescape(m.group(1) if m.group(1) is not None else m.group(2))


def decode_id(s):
    """Illustrator'ın id içindeki _x5F_ gibi kaçışlarını çözer."""
    return re.sub(r"_x([0-9A-Fa-f]{2,4})_", lambda m: chr(int(m.group(1), 16)), s)


def element_label(attrs):
    dn = get_attr(attrs, "data-name")
    if dn:
        return dn.strip()
    i = get_attr(attrs, "id")
    return decode_id(i).strip() if i else None


def parse_css(src):
    css = {}
    for sm in re.finditer(r"<style[^>]*>(.*?)</style>", src, re.S):
        body = re.sub(r"/\*.*?\*/", "", sm.group(1), flags=re.S)
        body = body.replace("<![CDATA[", "").replace("]]>", "")
        for rm in re.finditer(r"([^{}]+)\{([^{}]*)\}", body):
            props = {}
            for d in rm.group(2).split(";"):
                k, _, v = d.partition(":")
                if v.strip():
                    props[k.strip().lower()] = v.strip()
            for sel in rm.group(1).split(","):
                sel = sel.strip()
                if re.fullmatch(r"\.[\w\-]+", sel):
                    css.setdefault(sel[1:], {}).update(props)
    return css


def css_values(attrs, css, prop):
    """Bir öğenin style özelliği, sunum özelliği ve sınıflarından prop değerleri (öncelik sırasıyla)."""
    cands = []
    sm = STYLE_ATTR_RE.search(attrs)
    if sm:
        for d in sm.group(2).split(";"):
            k, _, v = d.partition(":")
            if k.strip().lower() == prop:
                cands.append(v)
    a = get_attr(attrs, prop)
    if a:
        cands.append(a)
    for c in (get_attr(attrs, "class") or "").split():
        v = css.get(c, {}).get(prop)
        if v:
            cands.append(v)
    return cands


def find_font_size(attrs, css):
    """(değer, birim) döndürür; bulunamazsa None."""
    for c in css_values(attrs, css, "font-size"):
        m = SIZE_RE.match(c)
        if m:
            try:
                return float(m.group(1)), m.group(2) or "px"
            except ValueError:
                pass
    return None


# ----------------------------------------------------------------------------
# Çekirdek: font karakter desteği (yazıdaki harfler fontta var mı?)
# ----------------------------------------------------------------------------

TR_CHARS = "çÇğĞıİöÖşŞüÜ"
FONT_MAGIC = (b"\x00\x01\x00\x00", b"OTTO", b"true", b"ttcf")
_font_index = None
_cmap_cache = {}


def font_dirs():
    env = os.environ
    return [
        (Path(env.get("WINDIR", r"C:\Windows")) / "Fonts", False),
        (Path(env.get("LOCALAPPDATA", "")) / "Microsoft" / "Windows" / "Fonts", False),
        # Adobe Fonts ile eşitlenen fontlar uzantısız, gizli dosyalardır.
        (Path(env.get("APPDATA", "")) / "Adobe" / "CoreSync" / "plugins" / "livetype", True),
        (Path(env.get("CommonProgramFiles", "")) / "Adobe" / "Fonts", True),
    ]


def _face_offsets(f):
    f.seek(0)
    head = f.read(12)
    if head[:4] not in FONT_MAGIC:
        return []
    if head[:4] == b"ttcf":
        n = struct.unpack(">I", head[8:12])[0]
        return list(struct.unpack(f">{n}I", f.read(4 * n)))
    return [0]


def _tables(f, off):
    f.seek(off + 4)
    n = struct.unpack(">H", f.read(2))[0]
    f.seek(off + 12)
    raw = f.read(16 * n)
    return {raw[i:i + 4]: struct.unpack(">II", raw[i + 8:i + 16]) for i in range(0, 16 * n, 16)}


def _font_names(f, off):
    tab = _tables(f, off).get(b"name")
    if not tab:
        return set()
    f.seek(tab[0])
    data = f.read(tab[1])
    count, so = struct.unpack(">HH", data[2:6])
    out = set()
    for i in range(count):
        pid, _eid, _lid, nid, ln, no = struct.unpack(">6H", data[6 + 12 * i:18 + 12 * i])
        if nid in (1, 4, 6, 16):  # aile, tam ad, PostScript adı, tipografik aile
            raw = data[so + no:so + no + ln]
            try:
                out.add(raw.decode("utf-16-be" if pid in (0, 3) else "mac_roman"))
            except UnicodeDecodeError:
                pass
    return out


def font_index():
    """Kurulu fontlar: {ad anahtarı: (dosya, yüz konumu)}. İlk çağrıda taranır."""
    global _font_index
    if _font_index is not None:
        return _font_index
    idx = {}
    for folder, any_ext in font_dirs():
        try:
            files = list(folder.rglob("*")) if any_ext else list(folder.glob("*"))
        except OSError:
            continue
        for p in files:
            if not any_ext and p.suffix.lower() not in (".ttf", ".otf", ".ttc", ".otc"):
                continue
            try:
                with open(p, "rb") as f:
                    for off in _face_offsets(f):
                        for name in _font_names(f, off):
                            idx.setdefault(norm(name), (p, off))
            except (OSError, struct.error, IndexError):
                continue
    _font_index = idx
    return idx


class FontFace:
    """Bir font yüzünün harf → glif eşlemesi ve glif genişlikleri (yalnızca gereken tablolar okunur)."""

    def __init__(self, path, off):
        with open(path, "rb") as f:
            tabs = _tables(f, off)

            def table(tag):
                if tag not in tabs:
                    return b""
                f.seek(tabs[tag][0])
                return f.read(tabs[tag][1])

            cmap, head, hhea, self.hmtx = table(b"cmap"), table(b"head"), table(b"hhea"), table(b"hmtx")
        self.upm = struct.unpack(">H", head[18:20])[0] or 1000
        self.nhm = struct.unpack(">H", hhea[34:36])[0]
        best = None
        for i in range(struct.unpack(">H", cmap[2:4])[0]):
            pid, eid, so = struct.unpack(">HHI", cmap[4 + 8 * i:12 + 8 * i])
            fmt = struct.unpack(">H", cmap[so:so + 2])[0]
            if fmt in (4, 12) and (pid == 0 or (pid == 3 and eid in (1, 10))):
                if best is None or fmt > best[0]:
                    best = (fmt, so)
        if best is None:
            raise ValueError("desteklenen cmap tablosu yok")
        self.fmt, self.sub = best[0], cmap[best[1]:]
        self._glyphs = {}

    def _u16(self, pos):
        return struct.unpack(">H", self.sub[pos:pos + 2])[0]

    def glyph(self, cp):
        """Kod noktasının glif numarası; fontta yoksa 0."""
        if cp not in self._glyphs:
            try:
                self._glyphs[cp] = self._lookup(cp)
            except struct.error:
                self._glyphs[cp] = 0
        return self._glyphs[cp]

    def _lookup(self, cp):
        s = self.sub
        if self.fmt == 12:
            for g in range(struct.unpack(">I", s[12:16])[0]):
                a, b, gid = struct.unpack(">III", s[16 + 12 * g:28 + 12 * g])
                if cp < a:
                    return 0
                if cp <= b:
                    return gid + cp - a
            return 0
        if cp > 0xFFFF:
            return 0
        seg = self._u16(6) // 2
        for i in range(seg):
            if cp <= self._u16(14 + 2 * i):
                start = self._u16(16 + 2 * seg + 2 * i)
                if cp < start:
                    return 0
                delta = self._u16(16 + 4 * seg + 2 * i)
                ro_pos = 16 + 6 * seg + 2 * i
                ro = self._u16(ro_pos)
                if ro == 0:
                    return (cp + delta) & 0xFFFF
                gid = self._u16(ro_pos + ro + 2 * (cp - start))
                return (gid + delta) & 0xFFFF if gid else 0
        return 0

    def width(self, text, size):
        """Yazının yaklaşık genişliği (harf aralığı ayarı/kerning hesaba katılmaz)."""
        total = 0
        for ch in text:
            i = min(self.glyph(ord(ch)), self.nhm - 1)
            total += struct.unpack(">H", self.hmtx[4 * i:4 * i + 2])[0]
        return total * size / self.upm


def find_face(family_css):
    """CSS font-family değerinden kurulu fontu bulur. (FontFace | None, gösterilecek ad)"""
    names = [x.strip().strip("\"'") for x in (family_css or "").split(",") if x.strip()]
    shown = names[-1] if names else "?"
    hit = next((font_index()[norm(n)] for n in names if norm(n) in font_index()), None)
    if hit is None:
        return None, shown
    key = (str(hit[0]), hit[1])
    if key not in _cmap_cache:
        try:
            _cmap_cache[key] = FontFace(*hit)
        except (OSError, struct.error, IndexError, ValueError):
            _cmap_cache[key] = None
    return _cmap_cache[key], shown


def font_missing(family_css, text):
    """Yazının fontta olmayan karakterleri. (durum, eksikler, font adı)
    durum: 'ok' | 'missing' | 'nofont' (font bu bilgisayarda yok, denetlenemedi)."""
    face, shown = find_face(family_css)
    if face is None:
        return "nofont", [], shown
    missing = [ch for ch in dict.fromkeys(text) if not ch.isspace() and not face.glyph(ord(ch))]
    return ("missing" if missing else "ok"), missing, shown


def font_checks(fname, family, val):
    """Elle kontrol gerektiren font sorunları (yazıda fontta olmayan harf)."""
    if not val:
        return []
    state, miss, shown = font_missing(family, val)
    if state == "missing":
        return [f"{fname}: '{shown}' fontunda olmayan karakter: {' '.join(miss)}"]
    if state == "nofont" and any(ord(ch) > 127 for ch in val):
        return [f"{fname}: '{shown}' fontu bu bilgisayarda bulunamadı, özel karakterler denetlenemedi: "
                f"{' '.join(dict.fromkeys(ch for ch in val if ord(ch) > 127))}"]
    return []


def field_limits(cfg):
    """(maks. karakter, en küçük font yüzdesi)"""
    try:
        mx = int(cfg.get("max") or 0)
    except (TypeError, ValueError):
        mx = 0
    try:
        minpct = float(cfg.get("minpct") or 0)
    except (TypeError, ValueError):
        minpct = 0
    return mx, minpct


def set_attr(tag, name, val):
    """Bir açılış etiketinde özelliği yazar (varsa değiştirir)."""
    m = re.search(r"(?<![\w:\-])%s\s*=\s*(?:\"[^\"]*\"|'[^']*')" % re.escape(name), tag)
    if m:
        return tag[:m.start()] + f'{name}="{val}"' + tag[m.end():]
    end = len(tag) - 2 if tag.endswith("/>") else len(tag) - 1
    return tag[:end].rstrip() + f' {name}="{val}"' + tag[end:]


def first_number(s, default=0.0):
    m = re.match(r"\s*(-?[\d.]+(?:e-?\d+)?)", s or "")
    try:
        return float(m.group(1)) if m else default
    except ValueError:
        return default


def set_style(tag, prop, val):
    """Bir açılış etiketinin style özelliğine prop: val yazar (varsa değiştirir)."""
    m = STYLE_ATTR_RE.search(tag)
    if m:
        decls = [d.strip() for d in m.group(2).split(";")
                 if d.strip() and d.split(":")[0].strip().lower() != prop]
        decls.append(f"{prop}: {val}")
        return tag[:m.start(2)] + "; ".join(decls) + tag[m.end(2):]
    end = len(tag) - 2 if tag.endswith("/>") else len(tag) - 1
    return tag[:end].rstrip() + f' style="{prop}: {val}"' + tag[end:]


def scan_texts(src):
    """Belgedeki <text> öğelerini konumlarıyla ve üst grubuyla birlikte bulur."""
    nodes, stack, pos = [], [], 0
    while True:
        m = TAG_RE.search(src, pos)
        if not m:
            break
        pos = m.end()
        name = m.group(2)
        if not name:
            continue
        attrs = m.group(3) or ""
        if m.group(1) == "/":
            for i in range(len(stack) - 1, -1, -1):
                if stack[i]["name"] == name:
                    del stack[i:]
                    break
            continue
        selfclose = attrs.rstrip().endswith("/")
        lname = name.split(":")[-1]
        if lname == "text" and not selfclose:
            cm = re.compile(r"</%s\s*>" % re.escape(name)).search(src, m.end())
            if not cm:
                continue
            if stack:
                stack[-1]["texts"] += 1
            nodes.append({"open_start": m.start(), "open_end": m.end(), "inner_end": cm.start(),
                          "attrs": attrs, "parent": stack[-1] if stack else None})
            pos = cm.end()
            continue
        if not selfclose:
            stack.append({"name": name, "lname": lname, "attrs": attrs, "texts": 0})
    return nodes


def split_field(label, stem):
    """'NAME_XL' + 'XL' -> ('NAME', True). Sonek şablon adıyla eşleşmezse (label, False)."""
    last = re.split(r"[_\-\s]+", stem)[-1]
    low = tr_lower(label)
    for cand in sorted({stem, last}, key=len, reverse=True):
        c = tr_lower(cand)
        if c and len(low) > len(c) + 1 and low.endswith(c) and low[-len(c) - 1] in "_- .":
            field = label[:len(label) - len(c)].rstrip("_- .")
            if field:
                return field, True
    return label, False


class Template:
    kind = "svg"

    def __init__(self, path):
        self.path = Path(path)
        self.stem = self.path.stem
        raw = self.path.read_bytes()
        if raw.startswith(b"\xef\xbb\xbf"):
            self.enc = "utf-8-sig"
        elif raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
            self.enc = "utf-16"
        else:
            m = re.search(rb"encoding\s*=\s*[\"']([\w\-]+)[\"']", raw[:200])
            self.enc = m.group(1).decode("ascii") if m else "utf-8"
        try:
            self.src = raw.decode(self.enc)
        except (UnicodeDecodeError, LookupError):
            self.enc = "utf-8"
            self.src = raw.decode("utf-8", errors="replace")
        self.css = parse_css(self.src)
        self.fields = {}
        for n in scan_texts(self.src):
            label = element_label(n["attrs"])
            own = bool(label)
            if not label:
                p = n["parent"]
                if p and p["lname"] == "g" and p["texts"] == 1:
                    label = element_label(p["attrs"])
            if not label:
                continue
            field, stripped = split_field(label, self.stem)
            inner = self.src[n["open_end"]:n["inner_end"]]
            tm = TSPAN_RE.search(inner)
            size = (find_font_size(tm.group(0), self.css) if tm else None) or find_font_size(n["attrs"], self.css)
            fam = (css_values(tm.group(0), self.css, "font-family") if tm else []) \
                or css_values(n["attrs"], self.css, "font-family")
            n.update(own=own, stripped=stripped, family=fam[0] if fam else "",
                     px=size[0] if size and size[1] == "px" else None,
                     sample=html.unescape(re.sub(r"<[^>]+>", "", inner)).strip(),
                     size=(fmt_num(size[0]) + size[1]) if size else "")
            self.fields.setdefault(field, []).append(n)

    def extent(self, n):
        """Şablondaki örnek yazının yerel koordinattaki (sol, sağ) kenarları; ölçülemezse None.
        Illustrator ortalı/sağa yaslı yazıyı da sol başlangıç noktasıyla kaydeder; gerçek ortayı
        bulmak için örnek yazının genişliği fonttan ölçülür."""
        face, _ = find_face(n["family"])
        if face is None or not n["px"]:
            return None
        cur = first_number(get_attr(n["attrs"], "x"))
        left, right = None, cur
        for tok in re.findall(r"<tspan\b[^>]*>|</tspan>|[^<]+", self.src[n["open_end"]:n["inner_end"]]):
            if tok.startswith("<tspan"):
                x = get_attr(tok, "x")
                if x is not None:
                    cur = first_number(x)
            elif not tok.startswith("<") and tok.strip():
                if left is None:
                    left = cur
                cur += face.width(html.unescape(tok), n["px"])
                right = max(right, cur)
        if left is None:
            return None
        anchor = (css_values(n["attrs"], self.css, "text-anchor") or ["start"])[0].strip()
        shift = {"middle": (right - left) / 2, "end": right - left}.get(anchor, 0)
        return left - shift, right - shift

    def render(self, values, cfgs):
        """values: {alan: yazı}. (bytes, küçültme notları, elle kontrol gerektiren nedenler) döndürür."""
        edits, notes, checks = [], [], []
        for fname, nodes in self.fields.items():
            if fname not in values:
                continue
            val = values[fname]
            cfg = cfgs.get(fname, {})
            mx, minpct = field_limits(cfg)
            for n in nodes:
                open_tag = self.src[n["open_start"]:n["open_end"]]
                inner = self.src[n["open_end"]:n["inner_end"]]
                tm = TSPAN_RE.search(inner)
                ts_tag = tm.group(0) if tm else None
                if ts_tag and ts_tag.endswith("/>"):
                    ts_tag = ts_tag[:-2].rstrip() + ">"
                checks += font_checks(fname, n["family"], val)
                if mx > 0 and len(val) > mx:
                    checks.append(f"{fname}: uzun yazı ({len(val)} karakter, sınır {mx})")
                    base = find_font_size(ts_tag, self.css) if ts_tag else None
                    on_tspan = base is not None
                    if base is None:
                        base = find_font_size(n["attrs"], self.css)
                    if base is None:
                        notes.append(f"{fname}: {len(val)} karakter (sınır {mx}) ama font boyutu "
                                     f"şablonda bulunamadı, küçültülemedi.")
                    else:
                        ratio = mx / len(val)
                        floor = minpct / 100.0
                        hit = ratio < floor
                        ratio = max(ratio, floor)
                        new = fmt_num(base[0] * ratio) + base[1]
                        open_tag = set_style(open_tag, "font-size", new)
                        if on_tspan:
                            ts_tag = set_style(ts_tag, "font-size", new)
                        notes.append(f"{fname}: {len(val)} karakter (sınır {mx}) → font "
                                     f"{fmt_num(base[0])}{base[1]} → {new}"
                                     + (" (alt sınıra takıldı, taşabilir)" if hit else ""))
                align = cfg.get("align", "left")
                if align in ("center", "right"):
                    ext = self.extent(n)
                    if ext is None:
                        checks.append(f"{fname}: hizalanamadı (font bu bilgisayarda kurulu değil ya da boyutu "
                                      f"okunamadı), sola yaslı bırakıldı")
                    else:
                        x = fmt_num((ext[0] + ext[1]) / 2 if align == "center" else ext[1])
                        open_tag = set_style(open_tag, "text-anchor", "middle" if align == "center" else "end")
                        if ts_tag and get_attr(ts_tag, "x") is not None:
                            ts_tag = set_attr(ts_tag, "x", x)
                        else:
                            open_tag = set_attr(open_tag, "x", x)
                esc = xml_escape(val)
                new_inner = (ts_tag + esc + "</tspan>") if ts_tag else esc
                edits.append((n["open_start"], n["inner_end"], open_tag + new_inner))
        out = self.src
        for s, e, rep in sorted(edits, reverse=True):
            out = out[:s] + rep + out[e:]
        # Dosyanın kodlamasında olmayan harfler (ör. iso-8859-1'de ı, ş, ğ) &#305; biçiminde yazılır.
        return out.encode(self.enc, errors="xmlcharrefreplace"), notes, list(dict.fromkeys(checks))

    def out_ext(self, opts):
        return ".svg"

    def write(self, data, target, opts):
        target.write_bytes(data)


# ----------------------------------------------------------------------------
# Çekirdek: Illustrator (.ai / .eps şablonlar)
# ----------------------------------------------------------------------------
# .ai/.eps şablonlar Illustrator'ın kendisiyle açılıp değiştirilir (COM + ExtendScript). Böylece paragraf
# hizası, kerning, font ve efektler Illustrator'da nasılsa öyle kalır; çıktı EPS, PDF ya da AI kaydedilir.

AI_EXTS = (".ai", ".eps")
OUT_FORMATS = {"eps": "EPS", "pdf": "PDF", "ai": "AI", "tif": "TIFF"}
TIF_COLORS = {"cmyk": "CMYK", "rgb": "RGB"}

AI_JS = r"""
function ser(v){
 if(v===null||v===undefined) return "null";
 var t=typeof v;
 if(t=="number") return isFinite(v)?String(v):"null";
 if(t=="boolean") return v?"true":"false";
 if(t=="string"){var s='"';for(var i=0;i<v.length;i++){var c=v.charAt(i),k=v.charCodeAt(i);
  if(c=='"'||c=='\\') s+='\\'+c; else if(k<32) s+='\\u'+('000'+k.toString(16)).slice(-4); else s+=c;}
  return s+'"';}
 if(v instanceof Array){var a=[];for(var j=0;j<v.length;j++) a.push(ser(v[j]));return "["+a.join(",")+"]";}
 var o=[];for(var key in v) if(v.hasOwnProperty(key)) o.push(ser(key)+":"+ser(v[key]));return "{"+o.join(",")+"}";
}
function samePath(a,b){return a.replace(/\//g,"\\").toLowerCase()==b.replace(/\//g,"\\").toLowerCase();}
function findDoc(p){
 for(var i=0;i<app.documents.length;i++){try{if(samePath(app.documents[i].fullName.fsName,p)) return app.documents[i];}catch(e){}}
 return null;
}
function frameLabel(tf){
 if(tf.name) return [tf.name,true];
 try{var p=tf.parent;
  if(p.typename=="GroupItem"&&p.pageItems.length==1&&p.name) return [p.name,false];
  if(p.typename=="Layer"&&p.pageItems.length==1&&p.layers.length==0) return [p.name,false];
 }catch(e){}
 return ["",false];
}
function bounds(tf){var b=tf.geometricBounds;return [b[0],b[1],b[2],b[3]];}
function analyzeDoc(doc){
 var out=[];
 for(var i=0;i<doc.textFrames.length;i++){
  var tf=doc.textFrames[i],lb=frameLabel(tf);
  if(!lb[0]) continue;
  var r={i:i,label:lb[0],own:lb[1],text:tf.contents,family:"",ps:"",size:0,just:"",
         kind:String(tf.kind).replace("TextType.",""),bounds:bounds(tf)};
  try{var ca=tf.textRange.characters[0].characterAttributes;r.size=ca.size;r.family=ca.textFont.family;r.ps=ca.textFont.name;}catch(e){}
  try{r.just=String(tf.paragraphs[0].paragraphAttributes.justification).replace("Justification.","");}catch(e){}
  out.push(r);
 }
 return out;
}
function setText(tf,e){
 var undo=[],x=tf;
 while(x&&x.typename!="Document"){try{if(x.locked){x.locked=false;undo.push(x);}}catch(er){} x=x.parent;}
 try{
  tf.contents=e.text;
  if(e.size>0&&e.text.length) tf.textRange.characterAttributes.size=e.size;
  if(e.kind=="POINTTEXT"&&e.text.length&&(e.align=="center"||e.align=="right")){
   var b=tf.geometricBounds,o=e.bounds;
   tf.translate(e.align=="center"?(o[0]+o[2])/2-(b[0]+b[2])/2:o[2]-b[2],0);
  }
 }finally{for(var k=0;k<undo.length;k++) undo[k].locked=true;}
}
function saveDoc(d,path,fmt){
 var f=new File(path),o;
 if(fmt=="tif"){
  // Dışa aktarma belgenin yolunu değiştirmez; belge şablon yolunda açık kalır, ama hiç kaydedilmez.
  // Çalışma yüzeyiyle kırpmak için çalışma yüzeyi aralığı verilir; Illustrator bu durumda dosya adına
  // "-01" gibi bir ek koyar. Bu yüzden geçici adla yazılır, Python asıl ada taşır.
  o=new ExportOptionsTIFF();o.resolution=DATA.dpi;o.lZWCompression=true;
  o.saveMultipleArtboards=true;o.artboardRange=String(d.artboards.getActiveArtboardIndex()+1);
  o.antiAliasing=AntiAliasingMethod.ARTOPTIMIZED;
  o.imageColorSpace=DATA.color=="rgb"?ImageColorSpace.RGB:ImageColorSpace.CMYK;
  var base="__gv_tmp_"+(new Date().getTime());
  d.exportFile(new File(f.parent.fsName+"/"+base+".tif"),ExportType.TIFF,o);
  var got=Folder(f.parent).getFiles(base+"*.tif");
  if(!got.length) throw new Error("TIFF dosyası oluşmadı: "+path);
  return got[0].fsName;
 }
 if(fmt=="pdf"){o=new PDFSaveOptions();o.preserveEditability=true;}
 else if(fmt=="ai"){o=new IllustratorSaveOptions();o.embedICCProfile=true;}
 else{o=new EPSSaveOptions();o.embedAllFonts=true;o.embedLinkedFiles=true;o.saveMultipleArtboards=false;}
 d.saveAs(f,o);
}
"""

AI_ANALYZE = r"""
var d=findDoc(DATA.path),mine=!d;
if(mine) d=app.open(new File(DATA.path));
try{return ser({frames:analyzeDoc(d)});}finally{if(mine) d.close(SaveOptions.DONOTSAVECHANGES);}
"""

AI_WRITE = r"""
var d=DATA.cur?findDoc(DATA.cur):null;
if(!d){
 if(findDoc(DATA.tpl)) throw new Error("Şablon Illustrator'da açık, kapatıp yeniden deneyin: "+DATA.tpl);
 d=app.open(new File(DATA.tpl));
}
for(var n=0;n<DATA.edits.length;n++){
 var e=DATA.edits[n];
 if(e.i>=d.textFrames.length||frameLabel(d.textFrames[e.i])[0]!=e.label)
  throw new Error("Şablon analizden sonra değişmiş, yeniden analiz edin ("+e.label+")");
 setText(d.textFrames[e.i],e);
}
var made=saveDoc(d,DATA.out,DATA.fmt);
return ser({cur:d.fullName.fsName,made:made||null});
"""

AI_TO_TIFF = r"""
var d=findDoc(DATA.src),mine=!d;
if(mine) d=app.open(new File(DATA.src));
try{var made=saveDoc(d,DATA.out,"tif");}finally{if(mine) d.close(SaveOptions.DONOTSAVECHANGES);}
return ser({made:made});
"""

AI_CLOSE = r"""
for(var n=0;n<DATA.paths.length;n++){var d=findDoc(DATA.paths[n]);if(d) d.close(SaveOptions.DONOTSAVECHANGES);}
return ser({});
"""


class IllustratorError(Exception):
    pass


class Illustrator:
    """Açık Illustrator'a COM ile bağlanır (gerekirse başlatır), ExtendScript çalıştırır.
    Üretimde şablon her satırda yeniden açılmaz: belge açık tutulur, farklı kaydet ile yeni dosyalar yazılır.
    Orijinal şablon hiç kaydedilmez."""

    def __init__(self):
        self.app = None
        self.docs = {}  # şablon yolu -> {"cur": açık belgenin şu anki yolu, "dirty": {çerçeve: geri alma düzenlemesi}}

    def call(self, body, **data):
        if self.app is None:
            try:
                import win32com.client
            except ImportError:
                raise IllustratorError(".ai/.eps şablonlar için pywin32 gerekli.\n\nKomut satırında:  pip install pywin32")
            try:
                self.app = win32com.client.Dispatch("Illustrator.Application")
            except Exception as e:
                raise IllustratorError(f"Illustrator'a bağlanılamadı (kurulu mu?): {e}")
        script = (AI_JS + "var DATA=" + json.dumps(data) + ";\n(function(){\n"
                  "var __ui=app.userInteractionLevel;app.userInteractionLevel=UserInteractionLevel.DONTDISPLAYALERTS;\n"
                  "try{" + body + "}catch(e){return ser({error:String(e.message||e)+(e.line?' (satır '+e.line+')':'')});}\n"
                  "finally{app.userInteractionLevel=__ui;}\n})();")
        try:
            res = self.app.DoJavaScript(script)
        except Exception as e:
            self.app = None  # Illustrator kapatılmış olabilir; sonraki çağrı yeniden bağlansın
            raise IllustratorError(f"Illustrator komutu çalışmadı: {e}")
        try:
            out = json.loads(res or "null")
        except ValueError:
            raise IllustratorError(f"Illustrator'dan beklenmeyen yanıt: {str(res)[:200]}")
        if isinstance(out, dict) and out.get("error"):
            raise IllustratorError(out["error"])
        return out

    def close_all(self):
        paths = [s["cur"] for s in self.docs.values() if s["cur"]]
        self.docs.clear()
        if paths:
            self.call(AI_CLOSE, paths=paths)


_illustrator = None
_ai_cache = None


def illustrator():
    global _illustrator
    if _illustrator is None:
        _illustrator = Illustrator()
    return _illustrator


def ai_cache_path():
    return session_path().with_name("grafik-varyant-ai-onbellek.json")


def ai_frames(path):
    """Şablondaki isimli yazı çerçeveleri. Illustrator'da açmak yavaş olduğu için sonuç, dosya değişene kadar
    saklanır; böylece program açılışında ve kontrolde Illustrator gerekmez."""
    global _ai_cache
    if _ai_cache is None:
        try:
            _ai_cache = json.loads(ai_cache_path().read_text(encoding="utf-8"))
        except (OSError, ValueError):
            _ai_cache = {}
    st = path.stat()
    key, sig = str(path.resolve()).lower(), [st.st_mtime_ns, st.st_size]
    hit = _ai_cache.get(key)
    if hit and hit.get("sig") == sig:
        return hit["frames"]
    frames = illustrator().call(AI_ANALYZE, path=str(path.resolve()))["frames"]
    _ai_cache[key] = {"sig": sig, "frames": frames}
    try:
        ai_cache_path().parent.mkdir(parents=True, exist_ok=True)
        ai_cache_path().write_text(json.dumps(_ai_cache, ensure_ascii=False), encoding="utf-8")
    except OSError:
        pass
    return frames


class AiTemplate:
    kind = "ai"

    def __init__(self, path):
        self.path = Path(path)
        self.stem = self.path.stem
        self.fields = {}
        self.nodes = {}
        for fr in ai_frames(self.path):
            field, stripped = split_field(fr["label"].strip(), self.stem)
            pt = float(fr.get("size") or 0)
            n = dict(fr, stripped=stripped, pt=pt,
                     family=", ".join(x for x in (fr.get("ps"), fr.get("family")) if x),
                     sample=re.sub(r"\s+", " ", fr.get("text", "")).strip(),
                     size=f"{fmt_num(pt)}pt" if pt else "")
            self.fields.setdefault(field, []).append(n)
            self.nodes[n["i"]] = n

    @staticmethod
    def _edit(n, text, size, align):
        return {"i": n["i"], "label": n["label"], "text": text.replace("\r\n", "\r").replace("\n", "\r"),
                "size": size, "align": align, "kind": n["kind"], "bounds": n["bounds"]}

    def render(self, values, cfgs):
        """(çerçeve düzenlemeleri, küçültme notları, elle kontrol gerektiren nedenler)"""
        edits, notes, checks = {}, [], []
        for fname, nodes in self.fields.items():
            if fname not in values:
                continue
            val = values[fname]
            cfg = cfgs.get(fname, {})
            mx, minpct = field_limits(cfg)
            for n in nodes:
                checks += font_checks(fname, n["family"], val)
                size = n["pt"]
                if mx > 0 and len(val) > mx:
                    checks.append(f"{fname}: uzun yazı ({len(val)} karakter, sınır {mx})")
                    if not size:
                        notes.append(f"{fname}: {len(val)} karakter (sınır {mx}) ama font boyutu "
                                     f"şablonda okunamadı, küçültülemedi.")
                    else:
                        floor = minpct / 100.0
                        ratio = mx / len(val)
                        size = n["pt"] * max(ratio, floor)
                        notes.append(f"{fname}: {len(val)} karakter (sınır {mx}) → font {fmt_num(n['pt'])}pt → "
                                     f"{fmt_num(size)}pt" + (" (alt sınıra takıldı, taşabilir)" if ratio < floor else ""))
                edits[n["i"]] = self._edit(n, val, size, cfg.get("align", "left"))
        return edits, notes, list(dict.fromkeys(checks))

    def out_ext(self, opts):
        return "." + (opts.get("fmt") or "eps")

    def write(self, edits, target, opts):
        ill = illustrator()
        key = str(self.path.resolve()).lower()
        state = ill.docs.setdefault(key, {"cur": None, "dirty": {}})
        # Belge önceki satırlardan açık kaldığı için, o satırlarda değişip bu satırda yazılmayan
        # çerçeveler şablondaki haline döndürülür.
        jobs = list(edits.values()) + [e for i, e in state["dirty"].items() if i not in edits]
        try:
            res = ill.call(AI_WRITE, tpl=str(self.path.resolve()), cur=state["cur"], out=str(target.resolve()),
                           fmt=opts.get("fmt") or "eps", edits=jobs, dpi=opts.get("tif_dpi") or 300,
                           color=opts.get("tif_color") or "cmyk")
        except IllustratorError:
            # Belgenin hangi durumda kaldığı belli değil; kapatılsın, sonraki satır şablonu yeniden açsın.
            cur = state["cur"]
            del ill.docs[key]
            if cur:
                try:
                    ill.call(AI_CLOSE, paths=[cur])
                except IllustratorError:
                    pass
            raise
        state["cur"] = res["cur"]
        take_made(res, target)
        state["dirty"] = {i: self._edit(self.nodes[i], self.nodes[i]["text"], self.nodes[i]["pt"], e["align"])
                          for i, e in edits.items()}


# ----------------------------------------------------------------------------
# Çekirdek: PDF şablonlar (Illustrator gerekmez)
# ----------------------------------------------------------------------------
# PDF'te yazı, içerik akışındaki "yazı göster" komutlarıdır (Tj, TJ, ', "). Alanın komutu yerinde
# değiştirilir: önceki renk, saydamlık, katman ve konum ayarları aynen geçerli kalır; yalnızca font
# bilgisayardaki tam fonta çevrilir (PDF'e gömülü font alt kümesinde yeni harfler bulunmayabilir).
# Alanlar iki yolla bulunur: yazının kendisi {NAME} biçiminde bir yer tutucuysa, ya da yazı NAME_M gibi
# adlandırılmış bir Acrobat katmanındaysa (Illustrator: PDF kaydederken "Üst düzey katmanlardan Acrobat
# katmanları oluştur").

PDF_WS = b" \t\r\n\f\x00"
PDF_DELIM = b"()<>[]{}/%"
PDF_SHOW = (b"Tj", b"TJ", b"'", b'"')
PLACEHOLDER_RE = re.compile(r"^\s*\{\s*([^{}]+?)\s*\}\s*$")


def _pdf_literal(raw):
    """(...) dizgesinin bayt değeri."""
    out, i, body = bytearray(), 0, raw[1:-1]
    esc = {ord("n"): 10, ord("r"): 13, ord("t"): 9, ord("b"): 8, ord("f"): 12}
    while i < len(body):
        c = body[i]
        if c != 0x5C:
            out.append(c)
            i += 1
            continue
        i += 1
        if i >= len(body):
            break
        c = body[i]
        if c in esc:
            out.append(esc[c])
            i += 1
        elif 0x30 <= c <= 0x37:
            j = i
            while j < len(body) and j < i + 3 and 0x30 <= body[j] <= 0x37:
                j += 1
            out.append(int(body[i:j], 8) & 0xFF)
            i = j
        elif c in (10, 13):  # satır devamı
            i += 2 if c == 13 and body[i + 1:i + 2] == b"\n" else 1
        else:
            out.append(c)
            i += 1
    return bytes(out)


def _pdf_hex(raw):
    h = re.sub(rb"[^0-9A-Fa-f]", b"", raw[1:-1])
    return bytes.fromhex((h + b"0" if len(h) % 2 else h).decode("ascii"))


def pdf_show_ops(data):
    """İçerik akışındaki yazı komutları ve form nesnesi çağrıları, sırasıyla.
    ("show", {...}) ya da ("do", ad). Her yazı komutu: bayt aralığı, o anki font kaynağı ve ayarlar, dizgeler."""
    out, stack, marks = [], [], []
    st = {"font": None, "tfs": 1.0, "tc": 0.0, "tw": 0.0, "th": 1.0, "ratio": 1.0, "bt": 0}
    saved = []
    i, n = 0, len(data)
    while i < n:
        c = data[i]
        if c in PDF_WS:
            i += 1
            continue
        if c == 0x25:  # yorum
            m = re.compile(rb"[\r\n]").search(data, i)
            i = m.end() if m else n
            continue
        start = i
        if c == 0x28:
            depth, j = 0, i
            while j < n:
                ch = data[j]
                if ch == 0x5C:
                    j += 2
                    continue
                if ch == 0x28:
                    depth += 1
                elif ch == 0x29:
                    depth -= 1
                    if depth == 0:
                        break
                j += 1
            stack.append(("str", _pdf_literal(data[i:j + 1]), start))
            i = j + 1
            continue
        if c == 0x3C:
            if data[i + 1:i + 2] == b"<":
                marks.append(len(stack))
                stack.append(("mark", None, start))
                i += 2
            else:
                j = data.find(b">", i)
                j = n - 1 if j < 0 else j
                stack.append(("str", _pdf_hex(data[i:j + 1]), start))
                i = j + 1
            continue
        if c == 0x3E and data[i + 1:i + 2] == b">":
            k = marks.pop() if marks else 0
            first = stack[k][2] if k < len(stack) else start
            del stack[k:]
            stack.append(("dict", None, first))
            i += 2
            continue
        if c == 0x5B:
            marks.append(len(stack))
            stack.append(("mark", None, start))
            i += 1
            continue
        if c == 0x5D:
            k = marks.pop() if marks else 0
            first = stack[k][2] if k < len(stack) else start
            items = [x for x in stack[k + 1:]]
            del stack[k:]
            stack.append(("arr", items, first))
            i += 1
            continue
        if c in (0x7B, 0x7D):
            i += 1
            continue
        j = i + 1
        while j < n and data[j] not in PDF_WS and data[j] not in PDF_DELIM:
            j += 1
        tok = data[i:j]
        i = j
        if c == 0x2F:
            stack.append(("name", tok[1:].decode("latin-1"), start))
            continue
        try:
            stack.append(("num", float(tok), start))
            continue
        except ValueError:
            pass
        if tok in (b"true", b"false", b"null"):
            stack.append(("kw", tok, start))
            continue
        # Operatör
        ops = stack
        stack, marks = [], []
        first = ops[0][2] if ops else start
        nums = [x[1] for x in ops if x[0] == "num"]
        if tok == b"ID":  # satır içi görsel: EI'ye kadar atla
            m = re.compile(rb"\sEI(?=[\s/\[<(]|$)").search(data, i)
            i = m.end() if m else n
        elif tok == b"q":
            saved.append(dict(st))
        elif tok == b"Q":
            if saved:
                bt = st["bt"]
                st = saved.pop()
                st["bt"] = bt
        elif tok == b"BT":
            st["bt"] += 1
            st["ratio"] = 1.0
        elif tok == b"Tf" and len(ops) >= 2 and ops[0][0] == "name":
            st["font"], st["tfs"] = ops[0][1], ops[1][1] if ops[1][0] == "num" else st["tfs"]
        elif tok == b"Tc" and nums:
            st["tc"] = nums[-1]
        elif tok == b"Tw" and nums:
            st["tw"] = nums[-1]
        elif tok == b"Tz" and nums:
            st["th"] = nums[-1] / 100.0
        elif tok == b"Tm" and len(nums) >= 4:
            a, b, cc, d = nums[:4]
            sy = (cc * cc + d * d) ** 0.5
            st["ratio"] = ((a * a + b * b) ** 0.5) / sy if sy else 1.0
        elif tok == b'"' and len(nums) >= 2:
            st["tw"], st["tc"] = nums[0], nums[1]
        elif tok == b"Do" and ops and ops[0][0] == "name":
            out.append(("do", ops[0][1]))
        if tok in PDF_SHOW:
            strings = []
            for x in ops:
                if x[0] == "str":
                    strings.append(x[1])
                elif x[0] == "arr":
                    strings += [y[1] for y in x[1] if y[0] == "str"]
            out.append(("show", dict(st, op=tok.decode(), start=first, end=i, strings=strings)))
    return out


def sfnt_face(path, off):
    """Font dosyasından tek yüz: .ttc koleksiyonundaysa ayrı bir .ttf/.otf olarak çıkarılır."""
    raw = Path(path).read_bytes()
    if raw[:4] != b"ttcf":
        return raw
    with open(path, "rb") as f:
        tabs = _tables(f, off)
    tags = sorted(tabs)
    num = len(tags)
    es = max(k for k in range(16) if 2 ** k <= num)
    head = raw[off:off + 4] + struct.pack(">HHHH", num, 16 * 2 ** es, es, num * 16 - 16 * 2 ** es)
    body, recs, pos = b"", b"", 12 + 16 * num
    for t in tags:
        o, ln = tabs[t]
        chunk = raw[o:o + ln]
        recs += t + struct.pack(">III", 0, pos + len(body), ln)
        body += chunk + b"\0" * (-ln % 4)
    return head + recs + body


def pdf_font_name(name):
    """'ABCDEF+Arial-BoldMT' -> 'Arial-BoldMT'"""
    return re.sub(r"^[A-Z]{6}\+", "", name or "")


def pdf_scan(doc):
    """Her yazı komutu için içerik akışındaki yeri, ayarları ve PyMuPDF'in okuduğu harfler (konum, font, katman)."""
    import pymupdf
    result = []
    for page in doc:
        fonts = {(f[6], f[4]): f[2] for f in page.get_fonts(full=True)}
        xobjs = {(x[2], x[1]): x[0] for x in page.get_xobjects()}
        ops = []

        def walk(xref, data, holder, depth=0):
            for kind, v in pdf_show_ops(data):
                if kind == "do":
                    x = xobjs.get((holder, v)) or xobjs.get((0, v))
                    if x and depth < 8:
                        walk(x, doc.xref_stream(x) or b"", x, depth + 1)
                    continue
                ftype = fonts.get((holder, v["font"])) or fonts.get((0, v["font"]), "")
                width = 2 if ftype == "Type0" else 1
                v.update(stream=xref, holder=holder, nchars=sum(len(s) // width for s in v["strings"]))
                ops.append(v)

        walk(0, page.read_contents(), 0)
        chars = [(ch, sp) for sp in page.get_texttrace() for ch in sp["chars"]]
        if sum(o["nchars"] for o in ops) != len(chars):
            raise ValueError(f"{page.number + 1}. sayfanın yazı yapısı çözümlenemedi "
                             f"({sum(o['nchars'] for o in ops)} / {len(chars)} harf)")
        pos = 0
        for o in ops:
            o["chars"] = chars[pos:pos + o["nchars"]]
            pos += o["nchars"]
            o["text"] = "".join(chr(ch[0]) for ch, _ in o["chars"] if ch[0] > 0)
        result.append(ops)
    return result


class PdfTemplate:
    kind = "pdf"

    def __init__(self, path):
        import pymupdf
        self.path = Path(path)
        self.stem = self.path.stem
        self.fields = {}
        with pymupdf.open(self.path) as doc:
            pages = pdf_scan(doc)
        for pno, ops in enumerate(pages):
            used = set()
            # 1) {ALAN} yer tutucuları: aynı BT bloğunda art arda gelen komutların birleşik yazısı
            for k, o in enumerate(ops):
                if k in used or not o["text"].strip().startswith("{"):
                    continue
                text, group = "", []
                for j in range(k, len(ops)):
                    if ops[j]["bt"] != o["bt"] or ops[j]["stream"] != o["stream"]:
                        break
                    text += ops[j]["text"]
                    group.append(j)
                    m = PLACEHOLDER_RE.match(text)
                    if m or "}" in text:
                        break
                m = PLACEHOLDER_RE.match(text)
                if m:
                    used.update(group)
                    self._add(m.group(1), True, False, pno, ops, group)
            # 2) ALAN_ŞABLON adlı katmandaki yazılar
            layers = {}
            for k, o in enumerate(ops):
                if k not in used and o["chars"] and o["chars"][0][1].get("layer"):
                    layers.setdefault(o["chars"][0][1]["layer"], []).append(k)
            for name, group in layers.items():
                field, stripped = split_field(name.strip(), self.stem)
                if stripped:
                    self._add(field, False, True, pno, ops, group)

    def _add(self, field, own, stripped, pno, ops, group):
        group = [k for k in group if ops[k]["chars"]] or group
        first, last = ops[group[0]], ops[group[-1]]
        sp = first["chars"][0][1] if first["chars"] else {"font": "", "size": 0, "dir": (1, 0)}
        dx, dy = sp.get("dir", (1, 0))
        if first["chars"] and last["chars"]:
            o0, b1 = first["chars"][0][0][2], last["chars"][-1][0][3]
            # Yazı yönünde baştaki harfin başlangıcından sondakinin bitişine kadar uzunluk
            right = max((b1[0] - o0[0]) * dx + (b1[1] - o0[1]) * dy, (b1[2] - o0[0]) * dx + (b1[3] - o0[1]) * dy)
        else:
            right = 0
        font = pdf_font_name(sp.get("font"))
        size = float(sp.get("size") or 0)
        text = "".join(ops[k]["text"] for k in group)
        self.fields.setdefault(field, []).append({
            "page": pno, "ops": [{key: ops[k][key] for key in
                                  ("stream", "holder", "start", "end", "op", "font", "tfs", "tc", "tw", "th", "ratio")}
                                 for k in group],
            "own": own, "stripped": stripped, "family": font, "pt": size, "width": right,
            "scale": size / first["tfs"] if first["tfs"] else 1.0,
            "sample": re.sub(r"\s+", " ", text).strip(), "size": f"{fmt_num(size)}pt" if size else ""})

    def render(self, values, cfgs):
        """(alan düzenlemeleri, küçültme notları, elle kontrol gerektiren nedenler)"""
        edits, notes, checks = [], [], []
        for fname, nodes in self.fields.items():
            if fname not in values:
                continue
            val = values[fname]
            cfg = cfgs.get(fname, {})
            mx, minpct = field_limits(cfg)
            for n in nodes:
                if find_face(n["family"])[0] is None:
                    checks.append(f"{fname}: '{n['family']}' fontu bu bilgisayarda yok, PDF'e yazılamaz")
                else:
                    checks += font_checks(fname, n["family"], val)
                ratio = 1.0
                if mx > 0 and len(val) > mx:
                    checks.append(f"{fname}: uzun yazı ({len(val)} karakter, sınır {mx})")
                    floor = minpct / 100.0
                    ratio = max(mx / len(val), floor)
                    notes.append(f"{fname}: {len(val)} karakter (sınır {mx}) → font {fmt_num(n['pt'])}pt → "
                                 f"{fmt_num(n['pt'] * ratio)}pt"
                                 + (" (alt sınıra takıldı, taşabilir)" if mx / len(val) < floor else ""))
                edits.append((n, val, ratio, cfg.get("align", "center")))
        return edits, notes, list(dict.fromkeys(checks))

    def out_ext(self, opts):
        return ".pdf"

    def write(self, edits, target, opts):
        import pymupdf
        doc = pymupdf.open(self.path)
        try:
            self._apply(doc, edits)
            for page in doc:
                # Illustrator'ın kendi düzenleme verisi kalırsa Illustrator dosyayı eski haliyle açar.
                doc.xref_set_key(page.xref, "PieceInfo", "null")
            try:
                doc.subset_fonts()  # fontTools gerekir; yoksa fontlar tam gömülür (dosya büyür)
            except Exception:
                pass
            doc.save(str(target), garbage=3, deflate=True)
        finally:
            doc.close()

    def _apply(self, doc, edits):
        import pymupdf
        splices = {}  # (sayfa, akış) -> [(baş, son, yeni bayt)]
        fonts = {}  # (sayfa, kaynak sahibi, font adı) -> (kaynak adı, pymupdf.Font)
        for n, val, ratio, align in edits:
            face, shown = find_face(n["family"])
            if face is None:
                raise OSError(f"'{shown}' fontu bu bilgisayarda bulunamadı, PDF'e yazılamadı")
            page = doc[n["page"]]
            first = n["ops"][0]
            key = (n["page"], first["holder"], n["family"])
            if key not in fonts:
                buf = sfnt_face(*font_index()[norm(n["family"])])
                res = f"SVF{len(fonts)}"
                fx = page.insert_font(fontname=res, fontbuffer=buf)
                if first["holder"]:
                    kind, val_ = doc.xref_get_key(first["holder"], "Resources/Font")
                    if kind == "xref":
                        doc.xref_set_key(int(val_.split()[0]), res, f"{fx} 0 R")
                    else:
                        doc.xref_set_key(first["holder"], f"Resources/Font/{res}", f"{fx} 0 R")
                fonts[key] = (res, pymupdf.Font(fontbuffer=buf))
            res, font = fonts[key]
            text = val.replace("\r", " ").replace("\n", " ")
            tfs = first["tfs"] * ratio
            codes = "".join(f"{font.has_glyph(ord(ch)):04X}" for ch in text)
            # Yeni yazının, komutun yazı uzayındaki genişliği (harf/kelime aralığı ve yatay ölçek dahil)
            w_new = (font.text_length(text, fontsize=tfs) + first["tc"] * len(text)
                     + first["tw"] * text.count(" ")) * first["th"]
            w_old = n["width"] / (n["scale"] * first["ratio"]) if n["scale"] else 0
            shift = {"center": (w_old - w_new) / 2, "right": w_old - w_new}.get(align, 0) / first["th"] if text else 0
            pre = {"'": "T* ", '"': "T* "}.get(first["op"], "")
            new = f" {pre}/{res} {tfs:.4f} Tf "
            if shift:
                new += f"{shift:.3f} 0 Td <{codes}> Tj {-shift:.3f} 0 Td "
            else:
                new += f"<{codes}> Tj "
            if first["font"]:
                new += f"/{first['font']} {fmt_num(first['tfs'])} Tf "
            splices.setdefault((n["page"], first["stream"]), []).append((first["start"], first["end"], new.encode()))
            for o in n["ops"][1:]:  # alanın diğer parçaları silinir
                splices.setdefault((n["page"], o["stream"]), []).append(
                    (o["start"], o["end"], b" T* " if o["op"] in ("'", '"') else b" "))
        for (pno, xref), items in splices.items():
            page = doc[pno]
            data = page.read_contents() if xref == 0 else doc.xref_stream(xref)
            for s, e, rep in sorted(items, reverse=True):
                data = data[:s] + rep + data[e:]
            if xref == 0:
                cont = page.get_contents()
                doc.update_stream(cont[0], data)
                if len(cont) > 1:
                    page.set_contents(cont[0])
            else:
                doc.update_stream(xref, data)


def take_made(res, target):
    """Illustrator'ın geçici adla yazdığı dosyayı (TIFF) asıl adına taşır."""
    made = (res or {}).get("made")
    if made and Path(made).resolve() != Path(target).resolve():
        os.replace(made, target)


def convert_tiff(folder, recursive=False, dpi=300, color="cmyk", skip_existing=True, progress=None):
    """Klasördeki .eps/.ai dosyalarını Illustrator'da açıp yanlarına aynı adlı .tif olarak dışa aktarır.
    Kaynak dosyalar kaydedilmez."""
    folder = Path(folder)
    found = folder.rglob("*") if recursive else folder.glob("*")
    files = sorted((p for p in found if p.suffix.lower() in AI_EXTS and p.is_file()), key=lambda p: str(p).lower())
    rep = {"done": [], "skipped": [], "errors": [], "total": len(files)}
    for k, src in enumerate(files, 1):
        if progress:
            progress(k, len(files), src.name)
        out = src.with_suffix(".tif")
        if skip_existing and out.exists():
            rep["skipped"].append(str(out))
            continue
        try:
            res = illustrator().call(AI_TO_TIFF, src=str(src.resolve()), out=str(out.resolve()), dpi=dpi, color=color)
            take_made(res, out)
        except (IllustratorError, OSError) as e:
            rep["errors"].append(f"{src.name}: {e}")
            continue
        rep["done"].append(str(out))
    return rep


MODES = {
    "svg": {"title": "SVG", "exts": (".svg",), "cls": Template},
    "eps": {"title": "EPS (Illustrator ile)", "exts": AI_EXTS, "cls": AiTemplate},
    "pdf": {"title": "PDF (Illustrator'süz)", "exts": (".pdf",), "cls": PdfTemplate},
}


def analyze(folder, recursive=False, exclude=None, progress=None, mode="svg"):
    folder = Path(folder)
    exts = MODES[mode]["exts"]
    found = folder.rglob("*") if recursive else folder.glob("*")
    paths = sorted((p for p in found if p.suffix.lower() in exts and p.is_file()), key=lambda p: str(p).lower())
    if exclude:
        ex = Path(exclude).resolve()
        if ex != folder.resolve():
            paths = [p for p in paths if ex not in p.resolve().parents]
    out = []
    for k, p in enumerate(paths, 1):
        if progress:
            progress(k, len(paths), p.name)
        out.append(MODES[mode]["cls"](p))
    return out


def summarize(templates):
    info = {}
    for t in templates:
        for name, nodes in t.fields.items():
            d = info.setdefault(name, {"files": [], "sample": "", "size": "", "auto": False, "tr": ""})
            if not d["tr"]:
                state, miss, shown = font_missing(nodes[0]["family"], TR_CHARS)
                d["tr"] = "font bilgisi yok" if not nodes[0]["family"] else {
                    "ok": f"var ({shown})", "missing": f"EKSİK: {''.join(miss)} ({shown})",
                    "nofont": f"font kurulu değil ({shown})"}[state]
            d["files"].append(t.stem)
            d["sample"] = d["sample"] or nodes[0]["sample"]
            d["size"] = d["size"] or nodes[0]["size"]
            d["auto"] = d["auto"] or any(x["own"] or x["stripped"] for x in nodes)
    return info


def build_lookup(templates):
    m, tokens = {}, {}
    for t in templates:
        m[norm(t.stem)] = t
        tokens.setdefault(norm(re.split(r"[_\-\s]+", t.stem)[-1]), []).append(t)
    for k, v in tokens.items():
        if len(v) == 1:
            m.setdefault(k, v[0])
    # 2XL yazılmışsa XXL.svg'yi (ve tersini) bul.
    for a, b in SIZE_ALIASES:
        if a in m:
            m.setdefault(b, m[a])
        elif b in m:
            m.setdefault(a, m[b])
    return m


def new_field(name, use=True, align="left"):
    return {"name": name, "use": use, "case": "asis", "max": 0, "minpct": 60,
            "empty": "blank", "prefix": "", "suffix": "", "align": align}


def apply_cfg(raw, cfg):
    v = str(raw or "").strip()
    c = cfg.get("case")
    if c == "upper":
        v = tr_upper(v)
    elif c == "lower":
        v = tr_lower(v)
    elif c == "title":
        v = tr_title(v)
    if v:
        v = (cfg.get("prefix") or "") + v + (cfg.get("suffix") or "")
    return v


# ----------------------------------------------------------------------------
# Çekirdek: dosya adı ve üretim
# ----------------------------------------------------------------------------

INVALID_RE = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
RESERVED = {"con", "prn", "aux", "nul"} | {f"com{i}" for i in range(1, 10)} | {f"lpt{i}" for i in range(1, 10)}


def clean_part(s):
    s = INVALID_RE.sub("_", str(s)).strip().rstrip(". ")
    if s.lower() in RESERVED:
        s = "_" + s
    return s


def name_parts(pattern, row, idx, tname):
    """Kalıptan klasör/dosya parçaları üretir. (parçalar, bilinmeyen yer tutucular)"""
    unknown = []
    keys = {norm(k): k for k in row}

    def rep(m):
        key, f = m.group(1).strip(), m.group(2)
        nk = norm(key)
        if nk in keys and keys[nk] not in (COL_T, COL_F):
            return INVALID_RE.sub("_", str(row.get(keys[nk], "") or "").strip())
        if nk in [norm(a) for a in ALIAS_T]:
            return INVALID_RE.sub("_", tname)
        if nk in ("sira", "index", "sirano"):
            return str(idx).zfill(int(f)) if f and f.isdigit() else str(idx)
        if nk in ("tarih", "date"):
            return datetime.date.today().isoformat()
        unknown.append(key)
        return ""

    text = re.sub(r"\{([^{}:]+)(?::([^{}]*))?\}", rep, pattern or "")
    parts = [p for p in (clean_part(x) for x in re.split(r"[/\\]+", text)) if p]
    return parts, unknown


def row_label(row, active, idx):
    for f in active:
        v = str(row.get(f["name"], "") or "").strip()
        if v:
            return v
    return f"satır {idx}"


OUT_EXT_RE = re.compile(r"\.(svg|eps|pdf|ai|tiff?)$", re.I)


def generate(opts, templates, fields, rows, dry=False, progress=None):
    try:
        return _generate(opts, templates, fields, rows, dry, progress)
    finally:
        if _illustrator is not None and _illustrator.docs:
            try:
                _illustrator.close_all()
            except IllustratorError:
                pass


def _generate(opts, templates, fields, rows, dry, progress):
    rep = {"done": [], "missing": [], "skipped": [], "warn": [], "shrink": [], "errors": [], "dry": dry,
           "check": [], "check_dir": bool(opts.get("check_dir")),
           "total": 0}
    lookup = build_lookup(templates)
    tpl_paths = {str(t.path.resolve()).lower() for t in templates}
    out = Path(opts["out_dir"])
    active = [f for f in fields if f.get("use")]
    cfgs = {f["name"]: f for f in active}
    used = set()
    live = [(i, r) for i, r in enumerate(rows, 1) if any(str(v or "").strip() for v in r.values())]
    rep["total"] = len(live)
    for k, (i, row) in enumerate(live, 1):
        if progress:
            progress(k, len(live))
        label = row_label(row, active, i)
        tv = str(row.get(COL_T, "") or "").strip()
        if not tv and len(templates) == 1:
            t = templates[0]
        else:
            t = lookup.get(norm(tv)) if tv else None
        if t is None:
            rep["missing"].append((i, tv or "(boş)", label))
            continue
        values = {}
        for f in active:
            nm = f["name"]
            val = apply_cfg(row.get(nm, ""), f)
            if nm not in t.fields:
                if val:
                    rep["warn"].append(f"Satır {i} ({label}): '{t.stem}' şablonunda {nm} alanı yok, yazılamadı.")
                continue
            if not val:
                if f.get("empty") == "keep":
                    rep["warn"].append(f"Satır {i} ({label}): {nm} boş, şablondaki yazı bırakıldı.")
                    continue
                rep["warn"].append(f"Satır {i} ({label}): {nm} boş yazıldı.")
            values[nm] = val
        try:
            job, notes, checks = t.render(values, cfgs)
        except Exception as e:  # tek bir bozuk şablon tüm işi durdurmasın
            rep["errors"].append(f"Satır {i} ({label}): şablon işlenemedi: {e}")
            continue
        for n in notes:
            rep["shrink"].append(f"Satır {i} ({label}) – {n}")

        custom = str(row.get(COL_F, "") or "").strip()
        if opts.get("use_name_col", True) and custom:
            custom = OUT_EXT_RE.sub("", custom)
            parts = [p for p in (clean_part(x) for x in re.split(r"[/\\]+", custom)) if p]
        else:
            parts, unknown = name_parts(opts.get("pattern", ""), row, i, t.stem)
            for u in unknown:
                rep["warn"].append(f"Satır {i} ({label}): dosya adı kalıbında bilinmeyen alan {{{u}}}.")
        if not parts:
            parts = [f"satir_{i}"]
        if opts.get("subfolder"):
            parts = [clean_part(t.stem)] + parts
        if checks and opts.get("check_dir"):
            parts = [CHECK_DIR] + parts
        folder = out.joinpath(*parts[:-1])
        stem = parts[-1]
        ext = t.out_ext(opts)
        target = folder / (stem + ext)

        n = 2
        while str(target.resolve()).lower() in used:
            target = folder / f"{stem}_{n}{ext}"
            n += 1
        if n > 2:
            rep["warn"].append(f"Satır {i} ({label}): aynı dosya adı bu listede birden fazla, "
                               f"'{target.name}' olarak kaydedildi.")
        key = str(target.resolve()).lower()
        if key in tpl_paths and opts.get("protect", True):
            rep["skipped"].append((i, label, str(target), "orijinal şablonun üzerine yazılacaktı (koruma açık)"))
            continue
        if target.exists() and key not in used:
            mode = opts.get("collision", "overwrite")
            if mode == "skip":
                rep["skipped"].append((i, label, str(target), "dosya zaten var"))
                used.add(key)
                continue
            if mode == "number":
                while target.exists() or str(target.resolve()).lower() in used:
                    target = folder / f"{stem}_{n}{ext}"
                    n += 1
                key = str(target.resolve()).lower()
        used.add(key)
        if not dry:
            try:
                target.parent.mkdir(parents=True, exist_ok=True)
                t.write(job, target, opts)
            except (OSError, IllustratorError) as e:
                rep["errors"].append(f"Satır {i} ({label}): yazılamadı: {e}")
                continue
        rep["done"].append((i, t.stem, label, str(target)))
        if checks:
            rep["check"].append((i, label, str(target), checks))
    return rep


def report_text(rep, out_dir=""):
    L = []
    title = "KONTROL (deneme, dosya yazılmadı)" if rep["dry"] else "ÜRETİM RAPORU"
    L.append(f"{title} – {datetime.datetime.now():%d.%m.%Y %H:%M}")
    if out_dir:
        L.append(f"Çıktı klasörü: {out_dir}")
    L.append("")
    verb = "Üretilecek" if rep["dry"] else "Üretilen"
    L.append(f"Toplam satır: {rep['total']}   {verb}: {len(rep['done'])}   "
             f"Şablonu eksik: {len(rep['missing'])}   Atlanan: {len(rep['skipped'])}   Hata: {len(rep['errors'])}"
             f"   Kontrol gerekli: {len(rep['check'])}")
    L.append("")
    if rep["check"]:
        where = (f"'{CHECK_DIR}' klasörüne ayrıldı" if rep["check_dir"] else "diğerleriyle aynı klasörde")
        L.append(f"ELLE KONTROL EDİLMESİ GEREKENLER ({len(rep['check'])} dosya, {where})")
        for i, label, p, reasons in rep["check"]:
            L.append(f"  Satır {i} ({label}) → {Path(p).name}")
            L += ["      " + r for r in reasons]
        L.append("")
    if rep["missing"]:
        L.append("ŞABLONU EKSİK OLANLAR (üretilemedi)")
        groups = {}
        for i, tv, label in rep["missing"]:
            groups.setdefault(tv, []).append(label)
        for tv, labels in groups.items():
            L.append(f"  {tv}  ({len(labels)} adet): {', '.join(labels)}")
        L.append("")
    if rep["errors"]:
        L.append("HATALAR")
        L += ["  " + e for e in rep["errors"]]
        L.append("")
    if rep["skipped"]:
        L.append("ATLANANLAR")
        L += [f"  Satır {i} ({label}): {why} → {p}" for i, label, p, why in rep["skipped"]]
        L.append("")
    if rep["shrink"]:
        L.append("FONT KÜÇÜLTÜLENLER")
        L += ["  " + s for s in rep["shrink"]]
        L.append("")
    if rep["warn"]:
        L.append("UYARILAR")
        L += ["  " + w for w in rep["warn"]]
        L.append("")
    if rep["done"]:
        L.append("DOSYALAR")
        groups = {}
        for i, stem, label, p in rep["done"]:
            groups.setdefault(stem, []).append(Path(p).name)
        for stem, names in groups.items():
            L.append(f"  {stem}  ({len(names)} adet)")
            L += ["    " + n for n in names]
    return "\n".join(L)


# ----------------------------------------------------------------------------
# Çekirdek: tablo verisi okuma (yapıştırma, CSV, Excel)
# ----------------------------------------------------------------------------

def parse_text(text, delim=None):
    lines = [ln for ln in text.replace("\r\n", "\n").replace("\r", "\n").split("\n") if ln.strip()]
    if not lines:
        return []
    if delim is None:
        counts = {d: sum(ln.count(d) for ln in lines[:20]) for d in ("\t", ";", "|", ",")}
        best = max(counts, key=counts.get)
        delim = best if counts[best] > 0 else "\x00"
    rows = [[c.strip() for c in r] for r in csv.reader(lines, delimiter=delim, quotechar='"')]
    return tidy_rows(rows)


def tidy_rows(rows):
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    keep = [j for j in range(width) if any(r[j] for r in rows)]
    return [[r[j] for j in keep] for r in rows]


def read_table_file(path):
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = []
        for r in ws.iter_rows(values_only=True):
            cells = []
            for v in r:
                if v is None:
                    cells.append("")
                elif isinstance(v, float) and v.is_integer():
                    cells.append(str(int(v)))
                else:
                    cells.append(str(v).strip())
            rows.append(cells)
        wb.close()
        return tidy_rows(rows)
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1254"):
        try:
            return parse_text(raw.decode(enc))
        except UnicodeDecodeError:
            continue
    return parse_text(raw.decode("latin-1"))


def match_target(header, targets):
    h = norm(header)
    if not h:
        return ""
    for t in targets:
        if norm(t) == h:
            return t
    if h in [norm(a) for a in ALIAS_T]:
        return COL_T
    if h in [norm(a) for a in ALIAS_F]:
        return COL_F
    for g in ALIAS_GROUPS:
        ng = [norm(x) for x in g]
        if h in ng:
            for t in targets:
                if norm(t) in ng:
                    return t
    return ""


def template_col(rows, lookup):
    """Şablon adlarıyla en çok eşleşen sütunun sırası; hiç eşleşme yoksa None.
    Eksik şablonlu satırlar olabileceği için tam eşleşme aranmaz."""
    if not rows or not lookup:
        return None
    ratios = []
    for j in range(len(rows[0])):
        vals = [r[j] for r in rows if r[j]]
        ratios.append(sum(1 for v in vals if norm(v) in lookup) / len(vals) if vals else 0)
    return ratios.index(max(ratios)) if max(ratios) > 0 else None


def guess_mapping(rows, targets, lookup):
    """(başlık var mı, [her sütun için hedef])"""
    if not rows:
        return False, []
    width = len(rows[0])
    head = [match_target(c, targets) for c in rows[0]]
    if any(head):
        seen, out = set(), []
        for h in head:
            out.append(h if h and h not in seen else "")
            seen.add(h)
        return True, out
    mapping = [""] * width
    tcol = template_col(rows, lookup)
    if tcol is None and width > 1:
        # Şablonla eşleşmese de en az çeşit içeren kısa sütun büyük ihtimalle beden sütunudur.
        best = min(range(width), key=lambda j: len({r[j] for r in rows}))
        if len({r[best] for r in rows}) < len(rows) * 0.6:
            tcol = best
    if tcol is not None:
        mapping[tcol] = COL_T
    rest = [t for t in targets if t not in (COL_T, COL_F)]
    for j in range(width):
        if j != tcol and rest:
            mapping[j] = rest.pop(0)
    return False, mapping


def write_xlsx(path, columns, rows, stems, fields):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Veri"
    ws.append(columns)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append([str(r.get(c, "") or "") for c in columns])
    for j, c in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = max(16, len(c) + 4)
        for cell in ws[get_column_letter(j)]:
            cell.number_format = "@"
    ws.freeze_panes = "A2"
    listing = ",".join(stems)
    if stems and COL_T in columns and len(listing) < 250 and not any("," in s or '"' in s for s in stems):
        dv = DataValidation(type="list", formula1=f'"{listing}"', allow_blank=True)
        dv.error = "Bu şablon klasörde yok."
        dv.showErrorMessage = False
        col = get_column_letter(columns.index(COL_T) + 1)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}2000")
    info = wb.create_sheet("Açıklama")
    lines = [
        f"{APP} – veri şablonu",
        "",
        "Her satır bir çıktı dosyasıdır. 'Veri' sayfasındaki başlıkları değiştirmeyin.",
        f"{COL_T}: Kullanılacak şablon dosyasının adı (uzantısız). Mevcut şablonlar: {', '.join(stems) or '-'}",
        f"{COL_F}: İsteğe bağlı. Doluysa çıktı dosyası bu adla kaydedilir; boşsa programdaki kalıp kullanılır.",
        "Diğer sütunlar şablondaki değişken alanlardır: " + (", ".join(fields) or "-"),
    ]
    for ln in lines:
        info.append([ln])
    info.column_dimensions["A"].width = 110
    info["A1"].font = Font(bold=True, size=13)
    wb.save(path)


# ----------------------------------------------------------------------------
# Arayüz
# ----------------------------------------------------------------------------

def resource_path(rel):
    """Kaynaktan çalışırken proje klasörü, exe içinde PyInstaller'ın açtığı geçici klasör."""
    return Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent)) / rel


def session_path(old=False):
    # Taşınabilir exe: ayarlar exe'nin yanında durur (klasör yazılabilirse), yoksa kullanıcı profilinde.
    # old=True: programın eski adı (SVG Varyant Üretici) döneminden kalan oturum dosyası.
    if getattr(sys, "frozen", False):
        exe_dir = Path(sys.executable).resolve().parent
        if os.access(exe_dir, os.W_OK):
            return exe_dir / ("svg-varyant-oturum.json" if old else "grafik-varyant-oturum.json")
    return Path(os.environ.get("APPDATA") or Path.home()) / ("SVGVaryantUretici" if old else "GrafikVaryantUretici") \
        / "oturum.json"


def open_path(p):
    try:
        os.startfile(str(p))
    except Exception as e:
        messagebox.showerror(APP, f"Açılamadı:\n{p}\n\n{e}")


class BulkDialog(tk.Toplevel):
    """Toplu giriş: her sütun ayrı bir kutu, kutuda satır başına bir değer.
    Aynı satır numarasındaki değerler tablonun aynı satırını oluşturur."""

    def __init__(self, app, focus_col=None, paste_text=""):
        super().__init__(app.winfo_toplevel())
        self.app = app
        self.title("Toplu giriş")
        self.transient(app.winfo_toplevel())
        self.cols = app.data_columns()
        self.geometry(f"{min(1500, max(760, 110 + 215 * len(self.cols)))}x660")
        self.texts, self.heads = {}, {}
        self._driver = None
        mono = ("Consolas", 11)

        ttk.Label(self, justify="left", foreground="#444", text=(
            "Her kutu bir sütundur: her satıra bir değer yazın ya da yapıştırın. Aynı numaralı satırlar tek bir "
            "ürün olur.\nExcel'den birden fazla sütunu birlikte yapıştırırsanız yan kutulara dağıtılır; beden "
            "sütunu kendiliğinden Beden kutusuna gider.")).pack(anchor="w", padx=10, pady=(10, 6))

        body = ttk.Frame(self)
        body.pack(fill="both", expand=True, padx=10)
        body.rowconfigure(1, weight=1)
        self.gutter = tk.Text(body, width=4, font=mono, bg="#ececec", fg="#888", relief="flat", takefocus=0,
                              state="disabled", cursor="arrow")
        self.gutter.grid(row=1, column=0, sticky="ns")
        for j, c in enumerate(self.cols, 1):
            body.columnconfigure(j, weight=1, uniform="col")
            h = ttk.Label(body, text=c, font=("Segoe UI", 10, "bold"))
            h.grid(row=0, column=j, sticky="w", padx=3, pady=(0, 3))
            t = tk.Text(body, width=18, wrap="none", undo=True, font=mono)
            t.grid(row=1, column=j, sticky="nsew", padx=2)
            t.insert("1.0", "\n".join(str(r.get(c, "") or "") for r in app.rows))
            t.edit_reset()
            t.configure(yscrollcommand=lambda a, b, src=t: self._on_scroll(src, a, b))
            t.bind("<<Paste>>", lambda e, col=c: self._paste(col))
            t.bind("<KeyRelease>", lambda e: self._counts())
            for ev in ("<Enter>", "<FocusIn>", "<Key>", "<MouseWheel>"):
                t.bind(ev, lambda e, src=t: self._drive(src), add="+")
            bf = ttk.Frame(body)
            bf.grid(row=2, column=j, sticky="w", padx=2, pady=(3, 0))
            ttk.Button(bf, text="Temizle", width=8, command=lambda col=c: self._set(col, [])).pack(side="left")
            ttk.Button(bf, text="Aşağı doldur", command=lambda col=c: self._fill(col)).pack(side="left", padx=3)
            self.texts[c], self.heads[c] = t, h
        self.sb = ttk.Scrollbar(body, orient="vertical", command=self._yview)
        self.sb.grid(row=1, column=len(self.cols) + 1, sticky="ns")

        bot = ttk.Frame(self)
        bot.pack(fill="x", padx=10, pady=10)
        self.info = ttk.Label(bot, text="")
        self.info.pack(side="left")
        ttk.Button(bot, text="Vazgeç", command=self.destroy).pack(side="right")
        ttk.Button(bot, text="Uygula", style="Big.TButton", command=self.apply).pack(side="right", padx=6)

        first = self.texts.get(focus_col) or self.texts[self.cols[1] if len(self.cols) > 2 else self.cols[0]]
        first.focus_set()
        first.mark_set("insert", "end")
        self._driver = first
        if paste_text.strip():
            self._paste(focus_col or self.cols[0], paste_text)
        self._counts()
        self.bind("<Escape>", lambda e: self.destroy())
        self.bind("<Control-Return>", lambda e: self.apply())
        self.grab_set()

    # ---- içerik ----
    def lines(self, col):
        out = [ln.strip() for ln in self.texts[col].get("1.0", "end-1c").split("\n")]
        while out and not out[-1]:
            out.pop()
        return out

    def _set(self, col, lines):
        t = self.texts[col]
        t.delete("1.0", "end")
        t.insert("1.0", "\n".join(lines))
        self._counts()

    def _fill(self, col):
        """Boş satırları bir üstteki değerle doldurur (en uzun sütunun boyuna kadar)."""
        n = max(len(self.lines(c)) for c in self.cols)
        cur = self.lines(col)
        cur += [""] * (n - len(cur))
        last = ""
        for i, v in enumerate(cur):
            if v:
                last = v
            else:
                cur[i] = last
        self._set(col, cur)

    def _paste(self, col, clip=None):
        """Çok sütunlu yapıştırmayı kutulara dağıtır; tek sütunlu yapıştırmayı Tk'ya bırakır."""
        given = clip is not None
        if not given:
            try:
                clip = self.clipboard_get()
            except tk.TclError:
                return "break"
        delim = "\t" if "\t" in clip else ";" if ";" in clip else None
        rows = parse_text(clip, delim) if delim else []
        if not rows or len(rows[0]) < 2:
            if rows or given:  # tek sütun: tırnak/ayırıcıdan arındırıp sütunun sonuna ekle
                new = [r[0] for r in rows] if rows else [ln.strip() for ln in clip.splitlines() if ln.strip()]
                self._set(col, self.lines(col) + new)
                return "break"
            self.after_idle(self._counts)
            return None
        tcol = template_col(rows, build_lookup(self.app.templates))
        targets = [c for c in self.cols[self.cols.index(col):] if tcol is None or c != COL_T]
        plan = {}
        for k in range(len(rows[0])):
            if k == tcol:
                plan[COL_T] = k
            elif targets:
                plan[targets.pop(0)] = k
        base = max(len(self.lines(c)) for c in plan)
        for c, k in plan.items():
            cur = self.lines(c)
            self._set(c, cur + [""] * (base - len(cur)) + [r[k] for r in rows])
        return "break"

    def _counts(self):
        counts = {c: len(self.lines(c)) for c in self.cols}
        n = max(counts.values())
        for c in self.cols:
            self.heads[c].config(text=f"{c}  ({counts[c]})")
        self.gutter.config(state="normal")
        self.gutter.delete("1.0", "end")
        self.gutter.insert("1.0", "\n".join(f"{i:>3}" for i in range(1, n + 40)))
        self.gutter.config(state="disabled")
        if self._driver is not None:
            self.gutter.yview(self._driver.index("@0,0"))
        filled = {v for v in counts.values() if v}
        if len(filled) > 1:
            self.info.config(foreground="#b00020", text="Satır sayıları farklı: " + ",  ".join(
                f"{c} {v}" for c, v in counts.items() if v) + "   (eksik kalanlar boş yazılır)")
        else:
            self.info.config(foreground="#444", text=f"{n} satır")

    # ---- birlikte kaydırma ----
    def _drive(self, src):
        self._driver = src

    def _on_scroll(self, src, a, b):
        # Yalnızca kullanıcının kaydırdığı kutu diğerlerini sürükler; kısa kutular uzunları geri çekmesin.
        if src is not self._driver:
            return
        self.sb.set(a, b)
        top = src.index("@0,0")
        for t in list(self.texts.values()) + [self.gutter]:
            if t is not src:
                t.yview(top)

    def _yview(self, *args):
        longest = max(self.texts.values(), key=lambda t: int(t.index("end-1c").split(".")[0]))
        self._driver = longest
        longest.yview(*args)

    def apply(self):
        data = {c: self.lines(c) for c in self.cols}
        n = max(len(v) for v in data.values())
        rows = []
        for i in range(n):
            r = dict(self.app.rows[i]) if i < len(self.app.rows) else {}
            for c in self.cols:
                r[c] = data[c][i] if i < len(data[c]) else ""
            if any(r[c] for c in self.cols):
                rows.append(r)
        self.app.rows = rows
        self.app.refresh_table()
        self.destroy()


class ImportDialog(tk.Toplevel):
    """Dosyadan okunan satırları tablo sütunlarına eşler."""

    def __init__(self, app, rows, title="Veri içe aktar"):
        super().__init__(app.winfo_toplevel())
        self.app = app
        self.title(title)
        self.geometry("900x560")
        self.transient(app.winfo_toplevel())
        self.raw = rows
        self.combos = []
        self.header = tk.BooleanVar(value=False)
        self.mode = tk.StringVar(value="append")

        opt = ttk.Frame(self)
        opt.pack(fill="x", padx=10, pady=4)
        ttk.Checkbutton(opt, text="İlk satır başlık (veri değil)", variable=self.header,
                        command=self.fill_preview).pack(side="left")
        ttk.Radiobutton(opt, text="Mevcut satırlara ekle", variable=self.mode, value="append").pack(side="left", padx=(20, 4))
        ttk.Radiobutton(opt, text="Tabloyu değiştir", variable=self.mode, value="replace").pack(side="left")

        self.mapf = ttk.LabelFrame(self, text="Sütun eşleştirme (her sütun hangi alana yazılsın?)")
        self.mapf.pack(fill="x", padx=10, pady=4)

        pf = ttk.LabelFrame(self, text="Önizleme")
        pf.pack(fill="both", expand=True, padx=10, pady=4)
        self.tree = ttk.Treeview(pf, show="headings", height=8)
        sb = ttk.Scrollbar(pf, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.tree.pack(fill="both", expand=True)

        bot = ttk.Frame(self)
        bot.pack(fill="x", padx=10, pady=10)
        self.info = ttk.Label(bot, text="")
        self.info.pack(side="left")
        ttk.Button(bot, text="İptal", command=self.destroy).pack(side="right")
        ttk.Button(bot, text="İçe aktar", command=self.do_import).pack(side="right", padx=6)

        self.setup_mapping()
        self.grab_set()

    def setup_mapping(self):
        for w in self.mapf.winfo_children():
            w.destroy()
        self.combos = []
        targets = self.app.data_columns()
        if not self.raw:
            ttk.Label(self.mapf, text="Henüz veri yok.").grid(padx=6, pady=6)
            self.fill_preview()
            return
        has_head, mapping = guess_mapping(self.raw, targets, build_lookup(self.app.templates))
        self.header.set(has_head)
        sample = self.raw[1] if has_head and len(self.raw) > 1 else self.raw[0]
        for j, tgt in enumerate(mapping):
            cell = ttk.Frame(self.mapf)
            cell.grid(row=j // 5, column=j % 5, padx=6, pady=4, sticky="w")
            ttk.Label(cell, text=f"Sütun {j + 1}  (örn: {sample[j][:14]})").pack(anchor="w")
            v = tk.StringVar(value=tgt or SKIP)
            cb = ttk.Combobox(cell, textvariable=v, values=[SKIP] + targets, state="readonly", width=20)
            cb.pack()
            cb.bind("<<ComboboxSelected>>", lambda e: self.fill_preview())
            self.combos.append(v)
        self.fill_preview()

    def mapped_rows(self):
        data = self.raw[1:] if self.header.get() else self.raw
        maps = [(j, v.get()) for j, v in enumerate(self.combos) if v.get() != SKIP]
        return [{tgt: r[j] for j, tgt in maps} for r in data]

    def fill_preview(self):
        self.tree.delete(*self.tree.get_children())
        cols = [v.get() for v in self.combos]
        ids = [f"c{j}" for j in range(len(cols))]
        self.tree["columns"] = ids
        for j, c in enumerate(cols):
            self.tree.heading(ids[j], text=f"{j + 1}: {c}")
            self.tree.column(ids[j], width=150, stretch=True)
        data = self.raw[1:] if self.header.get() else self.raw
        for r in data[:200]:
            self.tree.insert("", "end", values=r)
        self.info.config(text=f"{len(data)} satır")

    def do_import(self):
        chosen = [v.get() for v in self.combos if v.get() != SKIP]
        if not chosen:
            messagebox.showwarning(APP, "En az bir sütunu bir alana eşleştirin.", parent=self)
            return
        if len(set(chosen)) != len(chosen):
            messagebox.showwarning(APP, "Aynı alan birden fazla sütuna atanmış.", parent=self)
            return
        rows = self.mapped_rows()
        if self.mode.get() == "replace":
            self.app.rows = rows
        else:
            self.app.rows.extend(rows)
        self.app.refresh_table()
        self.destroy()


MODE_HINT = {
    "svg": "SVG şablonlar: değişecek yazılara Illustrator'da ALAN_ŞABLON adı verin (ör. M.svg içinde NAME_M).",
    "eps": ".ai / .eps şablonlar Illustrator ile açılıp düzenlenir (Illustrator kurulu olmalı). Yazı nesnesine "
           "ya da yalnız onu içeren katmana ALAN_ŞABLON adı verin (ör. M.ai içinde NAME_M).",
    "pdf": "PDF şablonlar Illustrator olmadan düzenlenir. Alanı bulmak için yazıyı {NAME} gibi süslü parantezle "
           "yazın ya da NAME_M adlı bir katmana koyup PDF'i 'Acrobat katmanları oluştur' seçeneğiyle kaydedin.",
}
ALIGN_HINT = {
    "svg": "Hizalama: Illustrator ortalı yazıyı SVG'ye sola yaslı kaydeder; 'Ortalı' seçilirse yeni yazı örnek "
           "yazının fonttan ölçülen ortasına göre yerleşir (font bu bilgisayarda kurulu olmalı).",
    "eps": "Hizalama: Illustrator'daki paragraf hizası aynen korunur; 'Ortalı'/'Sağa yaslı' ayrıca yeni yazıyı "
           "örnek yazının ortasına/sağ kenarına oturtur.",
    "pdf": "Hizalama: PDF'te paragraf hizası bilgisi yoktur; yeni yazı örnek yazının sol kenarına, ortasına ya da "
           "sağ kenarına göre yerleşir. Yazı, PDF'teki fontun bu bilgisayarda kurulu tam haliyle yazılır.",
}


class App(tk.Tk):
    """Ana pencere: her şablon türü (SVG, EPS, PDF) ayrı bir sekmede, kendi ayarları ve tablosuyla."""

    def __init__(self, load_session=True):
        super().__init__()
        self.title(APP)
        try:
            self.iconbitmap(default=str(resource_path("assets/app.ico")))
        except tk.TclError:
            pass
        self.geometry("1180x800")
        self.minsize(980, 640)
        lh = tkfont.nametofont("TkDefaultFont").metrics("linespace")
        st = ttk.Style(self)
        st.configure("Treeview", rowheight=lh + 8)
        st.configure("Big.TButton", font=("Segoe UI", 11, "bold"), padding=8)
        st.configure("Mode.TNotebook.Tab", font=("Segoe UI", 11, "bold"), padding=(16, 6))

        self._menu()
        self.status = ttk.Label(self, text="Hazır", anchor="w", relief="sunken")
        self.status.pack(fill="x", side="bottom")
        self.tabs = ttk.Notebook(self, style="Mode.TNotebook")
        self.tabs.pack(fill="both", expand=True, padx=6, pady=6)
        self.spaces = {}
        for mode, info in MODES.items():
            w = Workspace(self.tabs, self, mode)
            self.tabs.add(w, text=info["title"])
            self.spaces[mode] = w
        self.tabs.bind("<<NotebookTabChanged>>", lambda e: self.update_title())
        self.protocol("WM_DELETE_WINDOW", self.on_close)
        if load_session:
            sp = session_path()
            if not sp.exists():
                sp = session_path(old=True)
            if sp.exists():
                try:
                    self.load_session(json.loads(sp.read_text(encoding="utf-8")))
                except Exception:
                    pass

    @property
    def ws(self):
        return self.nametowidget(self.tabs.select())

    def load_session(self, d):
        if "modes" not in d:  # eski oturum: tek sekme vardı (SVG)
            d = {"modes": {"svg": d}, "active": "svg"}
        for mode, w in self.spaces.items():
            if mode in d["modes"]:
                w.load_project_dict(d["modes"][mode])
        if d.get("active") in self.spaces:
            self.tabs.select(self.spaces[d["active"]])

    # ---- menü ve proje dosyası ----
    def _menu(self):
        mb = tk.Menu(self)
        m = tk.Menu(mb, tearoff=0)
        m.add_command(label="Yeni proje (bu sekme)", command=lambda: self.ws.new_project())
        m.add_command(label="Proje aç…", command=self.open_project, accelerator="Ctrl+O")
        m.add_command(label="Projeyi kaydet", command=self.save_project, accelerator="Ctrl+S")
        m.add_command(label="Projeyi farklı kaydet…", command=lambda: self.save_project(True))
        m.add_separator()
        m.add_command(label="Çıkış", command=self.on_close)
        mb.add_cascade(label="Dosya", menu=m)
        h = tk.Menu(mb, tearoff=0)
        h.add_command(label="Nasıl kullanılır?", command=self.show_help)
        mb.add_cascade(label="Yardım", menu=h)
        self.config(menu=mb)
        self.bind("<Control-s>", lambda e: self.save_project())
        self.bind("<Control-o>", lambda e: self.open_project())

    def update_title(self):
        try:
            pf = self.ws.project_file
        except (KeyError, AttributeError, tk.TclError):  # sekmeler henüz kuruluyor
            pf = None
        self.title(f"{APP} – {Path(pf).name}" if pf else APP)

    def open_project(self):
        p = filedialog.askopenfilename(title="Proje aç", filetypes=[("Varyant projesi", "*.svgproje"), ("Tümü", "*.*")])
        if not p:
            return
        try:
            d = json.loads(Path(p).read_text(encoding="utf-8"))
        except Exception as e:
            messagebox.showerror(APP, f"Proje okunamadı:\n{e}")
            return
        d["project_file"] = p
        w = self.spaces.get(d.get("mode"), self.spaces["svg"])
        self.tabs.select(w)
        w.load_project_dict(d)

    def save_project(self, ask=False):
        w = self.ws
        p = w.project_file
        if ask or not p:
            p = filedialog.asksaveasfilename(title="Projeyi kaydet", defaultextension=".svgproje",
                                             filetypes=[("Varyant projesi", "*.svgproje")])
            if not p:
                return
        w.project_file = p
        try:
            Path(p).write_text(json.dumps(w.project_dict(), ensure_ascii=False, indent=1), encoding="utf-8")
        except OSError as e:
            messagebox.showerror(APP, f"Kaydedilemedi:\n{e}")
            return
        self.update_title()
        self.status.config(text=f"Proje kaydedildi: {p}")

    def on_close(self):
        try:
            sp = session_path()
            sp.parent.mkdir(parents=True, exist_ok=True)
            d = {"version": 2, "active": self.ws.mode,
                 "modes": {m: w.project_dict() for m, w in self.spaces.items()}}
            sp.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
        except Exception:
            pass
        self.destroy()

    def show_help(self):
        messagebox.showinfo(APP, (
            "Üstteki sekmeler şablon türleridir; her sekmenin klasörleri, alanları ve tablosu ayrıdır:\n"
            "  • SVG: .svg şablonlardan SVG üretir.\n"
            "  • EPS (Illustrator ile): .ai/.eps şablonları Illustrator'da doldurup EPS, PDF ya da AI kaydeder.\n"
            "  • PDF (Illustrator'süz): .pdf şablonları doğrudan doldurup PDF kaydeder.\n\n"
            "1) Değişecek yazılara ALAN_ŞABLON adı verin (ör. M şablonunda NAME_M). PDF'te yazıyı {NAME} gibi\n"
            "    süslü parantezle de yazabilirsiniz. Dosya adı şablonun (bedenin) adıdır.\n\n"
            "2) 'Proje ve Analiz' sekmesinde şablon klasörünü seçip 'Analiz Et'e basın.\n"
            "    Bulunan alanlardan hangilerinin değişken olduğunu işaretleyin.\n\n"
            "3) 'Veri' sekmesinde tabloyu doldurun: 'Toplu giriş' ile sütun sütun yapıştırarak,\n"
            "    elle ya da Excel/CSV yükleyerek.\n\n"
            "4) 'İsimlendirme ve Seçenekler' sekmesinde dosya adı kalıbını ve kuralları belirleyin.\n\n"
            "5) 'Üret' sekmesinde önce 'Kontrol Et', sonra 'ÜRET'."))


class Workspace(ttk.Frame):
    """Bir şablon türünün tam iş akışı: klasörler ve analiz, veri, isimlendirme, üretim."""

    def __init__(self, master, app, mode):
        super().__init__(master)
        self.app, self.mode = app, mode
        self.status = app.status
        self.templates = []
        self.fields = []
        self.rows = []
        self.project_file = None
        self.last_report = ""
        self._loading = False
        self._editor = None

        self.v_tpl = tk.StringVar()
        self.v_out = tk.StringVar()
        self.v_rec = tk.BooleanVar(value=False)
        self.v_pattern = tk.StringVar(value="{Beden}_{NAME}")
        self.v_usename = tk.BooleanVar(value=True)
        self.v_sub = tk.BooleanVar(value=False)
        self.v_coll = tk.StringVar(value="overwrite")
        self.v_protect = tk.BooleanVar(value=True)
        self.v_savereport = tk.BooleanVar(value=True)
        self.v_openafter = tk.BooleanVar(value=False)
        self.v_checkdir = tk.BooleanVar(value=True)
        self.v_fmt = tk.StringVar(value="eps")
        self.v_tif_dpi = tk.StringVar(value="300")
        self.v_tif_color = tk.StringVar(value=TIF_COLORS["cmyk"])

        self.nb = ttk.Notebook(self)
        self.nb.pack(fill="both", expand=True, padx=4, pady=4)
        self._tab_project()
        self._tab_data()
        self._tab_options()
        self._tab_run()

        self.v_pattern.trace_add("write", lambda *a: self.update_name_preview())
        self.v_fmt.trace_add("write", lambda *a: self.update_name_preview())
        self.refresh_table()

    def project_dict(self):
        return {"version": 1, "mode": self.mode, "tpl_dir": self.v_tpl.get(), "out_dir": self.v_out.get(),
                "recursive": self.v_rec.get(), "pattern": self.v_pattern.get(),
                "use_name_col": self.v_usename.get(), "subfolder": self.v_sub.get(),
                "collision": self.v_coll.get(), "protect": self.v_protect.get(),
                "save_report": self.v_savereport.get(), "open_after": self.v_openafter.get(),
                "check_dir": self.v_checkdir.get(), "out_fmt": self.v_fmt.get(),
                "tif_dpi": self.tif_dpi(), "tif_color": self.tif_color(),
                "fields": self.fields, "rows": self.rows, "project_file": self.project_file}

    def load_project_dict(self, d):
        self.v_tpl.set(d.get("tpl_dir", ""))
        self.v_out.set(d.get("out_dir", ""))
        self.v_rec.set(bool(d.get("recursive")))
        self.v_pattern.set(d.get("pattern", "{Beden}_{NAME}"))
        self.v_usename.set(d.get("use_name_col", True))
        self.v_sub.set(bool(d.get("subfolder")))
        self.v_coll.set(d.get("collision", "overwrite") if d.get("collision") in COLLISION else "overwrite")
        self.v_protect.set(d.get("protect", True))
        self.v_savereport.set(d.get("save_report", True))
        self.v_openafter.set(bool(d.get("open_after")))
        self.v_checkdir.set(d.get("check_dir", True))
        self.v_fmt.set(d.get("out_fmt") if d.get("out_fmt") in OUT_FORMATS else "eps")
        self.v_tif_dpi.set(str(d.get("tif_dpi") or 300))
        self.v_tif_color.set(TIF_COLORS.get(d.get("tif_color"), TIF_COLORS["cmyk"]))
        self.fields = [dict(new_field(f.get("name", "")), **f) for f in d.get("fields", []) if f.get("name")]
        self.rows = [dict(r) for r in d.get("rows", []) if isinstance(r, dict)]
        for r in self.rows:  # eski projelerde sütunun adı "Şablon"du
            if COL_T_OLD in r:
                r.setdefault(COL_T, r.pop(COL_T_OLD))
        self.project_file = d.get("project_file")
        self.templates = []
        if self.v_tpl.get() and Path(self.v_tpl.get()).is_dir():
            self.run_analysis(silent=True)
        else:
            self.refresh_fields()
        self.refresh_table()
        self.app.update_title()

    def new_project(self):
        if not messagebox.askyesno(APP, f"'{MODES[self.mode]['title']}' sekmesindeki klasörler, alan ayarları ve "
                                        f"veri tablosu sıfırlansın mı?\n\n"
                                        "Diskteki şablon ve çıktı dosyalarına dokunulmaz.", icon="warning"):
            return
        self.load_project_dict({})

    # ---- sekme 1: proje ve analiz ----
    def _tab_project(self):
        tab = ttk.Frame(self.nb, padding=10)
        self.nb.add(tab, text=" 1. Proje ve Analiz ")
        g = ttk.LabelFrame(tab, text="Klasörler", padding=8)
        g.pack(fill="x")
        g.columnconfigure(1, weight=1)
        ttk.Label(g, text="Şablon (okuma) klasörü:").grid(row=0, column=0, sticky="w")
        ttk.Entry(g, textvariable=self.v_tpl).grid(row=0, column=1, sticky="ew", padx=6, pady=3)
        ttk.Button(g, text="Gözat…", command=lambda: self.browse(self.v_tpl, True)).grid(row=0, column=2)
        ttk.Button(g, text="Aç", width=5, command=lambda: self.open_dir(self.v_tpl)).grid(row=0, column=3, padx=(4, 0))
        ttk.Label(g, text="Çıktı klasörü:").grid(row=1, column=0, sticky="w")
        ttk.Entry(g, textvariable=self.v_out).grid(row=1, column=1, sticky="ew", padx=6, pady=3)
        ttk.Button(g, text="Gözat…", command=lambda: self.browse(self.v_out)).grid(row=1, column=2)
        ttk.Button(g, text="Aç", width=5, command=lambda: self.open_dir(self.v_out)).grid(row=1, column=3, padx=(4, 0))
        bar = ttk.Frame(g)
        bar.grid(row=2, column=0, columnspan=2, sticky="w", pady=(6, 0))
        ttk.Button(g, text="Her şeyi sıfırla…", command=self.new_project).grid(
            row=2, column=2, columnspan=2, sticky="e", pady=(6, 0))
        ttk.Button(bar, text="Analiz Et", style="Big.TButton", command=self.run_analysis).pack(side="left")
        ttk.Checkbutton(bar, text="Alt klasörleri de tara", variable=self.v_rec).pack(side="left", padx=12)
        self.lbl_tpl = ttk.Label(bar, text="Henüz analiz yapılmadı.", foreground="#555")
        self.lbl_tpl.pack(side="left", padx=8)
        ttk.Label(g, text=MODE_HINT[self.mode], foreground="#555", wraplength=1050, justify="left").grid(
            row=3, column=0, columnspan=4, sticky="w", pady=(8, 0))

        f = ttk.LabelFrame(tab, text="Bulunan alanlar (değişken olacakları 'Kullan' sütununa tıklayarak işaretleyin)",
                           padding=8)
        f.pack(fill="both", expand=True, pady=(10, 0))
        cols = ("use", "name", "files", "miss", "sample", "size", "tr", "align", "case", "max", "minpct", "empty",
                "prefix", "suffix")
        heads = ("Kullan", "Alan", "Bulunduğu şablonlar", "Eksik olduğu şablonlar", "Örnek yazı", "Font",
                 "Fontta Türkçe harf", "Hiza", "Harf", "Maks. karakter", "Min. font %", "Boşsa", "Önek", "Sonek")
        widths = (55, 110, 130, 120, 95, 75, 180, 80, 105, 95, 80, 130, 50, 50)
        self.ftree = ttk.Treeview(f, columns=cols, show="headings", height=7, selectmode="browse")
        for c, h, w in zip(cols, heads, widths):
            self.ftree.heading(c, text=h)
            self.ftree.column(c, width=w, anchor="center" if c in ("use", "max", "minpct", "size") else "w")
        sb = ttk.Scrollbar(f, orient="vertical", command=self.ftree.yview)
        self.ftree.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.ftree.pack(fill="both", expand=True)
        self.ftree.bind("<Button-1>", self._field_click)
        self.ftree.bind("<<TreeviewSelect>>", self._field_select)

        e = ttk.LabelFrame(tab, text="Seçili alanın ayarları", padding=8)
        e.pack(fill="x", pady=(10, 0))
        self.fv = {k: tk.StringVar() for k in ("case", "max", "minpct", "empty", "prefix", "suffix", "align")}
        self.lbl_field = ttk.Label(e, text="(alan seçin)", font=("Segoe UI", 10, "bold"))
        self.lbl_field.grid(row=0, column=0, columnspan=8, sticky="w", pady=(0, 6))
        ttk.Label(e, text="Harf dönüşümü:").grid(row=1, column=0, sticky="w")
        ttk.Combobox(e, textvariable=self.fv["case"], values=list(CASES.values()), state="readonly",
                     width=20).grid(row=1, column=1, padx=6, sticky="w")
        ttk.Label(e, text="Boş gelirse:").grid(row=1, column=2, sticky="w", padx=(14, 0))
        ttk.Combobox(e, textvariable=self.fv["empty"], values=list(EMPTY.values()), state="readonly",
                     width=24).grid(row=1, column=3, padx=6, sticky="w")
        ttk.Label(e, text="Önek:").grid(row=1, column=4, sticky="w", padx=(14, 0))
        ttk.Entry(e, textvariable=self.fv["prefix"], width=8).grid(row=1, column=5, padx=6)
        ttk.Label(e, text="Sonek:").grid(row=1, column=6, sticky="w")
        ttk.Entry(e, textvariable=self.fv["suffix"], width=8).grid(row=1, column=7, padx=6)
        ttk.Label(e, text="Maks. karakter:").grid(row=2, column=0, sticky="w", pady=(8, 0))
        ttk.Spinbox(e, textvariable=self.fv["max"], from_=0, to=200, width=6).grid(row=2, column=1, padx=6, sticky="w", pady=(8, 0))
        ttk.Label(e, text="En fazla şu kadar küçült (%):").grid(row=2, column=2, sticky="w", padx=(14, 0), pady=(8, 0))
        ttk.Spinbox(e, textvariable=self.fv["minpct"], from_=10, to=100, increment=5, width=6).grid(
            row=2, column=3, padx=6, sticky="w", pady=(8, 0))
        ttk.Label(e, text="Hizalama:").grid(row=2, column=4, sticky="w", padx=(14, 0), pady=(8, 0))
        ttk.Combobox(e, textvariable=self.fv["align"], values=list(ALIGN.values()), state="readonly",
                     width=14).grid(row=2, column=5, columnspan=3, padx=6, sticky="w", pady=(8, 0))
        ttk.Label(e, foreground="#555", justify="left", text=(
            "Yazı 'Maks. karakter'den uzunsa font boyutu orantılı küçültülür (0 = kapalı). "
            "Örn. sınır 8, yazı 10 harf → font %80. Alt sınır: orijinal boyutun yüzdesi.\n"
            + ALIGN_HINT[self.mode])
                  ).grid(row=3, column=0, columnspan=8, sticky="w", pady=(8, 0))
        for v in self.fv.values():
            v.trace_add("write", self._field_changed)

    def browse(self, var, analyze_after=False):
        d = filedialog.askdirectory(initialdir=var.get() or None)
        if not d:
            return
        var.set(os.path.normpath(d))
        if analyze_after:
            if not self.v_out.get():
                self.v_out.set(str(Path(var.get()).parent / "export"))
            self.run_analysis()

    def open_dir(self, var):
        p = var.get().strip()
        if p and Path(p).is_dir():
            open_path(p)
        else:
            messagebox.showinfo(APP, "Klasör henüz yok.")

    def run_analysis(self, silent=False):
        folder = self.v_tpl.get().strip()
        if not folder or not Path(folder).is_dir():
            if not silent:
                messagebox.showwarning(APP, "Geçerli bir şablon klasörü seçin.")
            return
        try:
            self.templates = analyze(folder, self.v_rec.get(), self.v_out.get().strip() or None,
                                     progress=self.analysis_progress, mode=self.mode)
        except Exception as e:
            self.status.config(text="Analiz yapılamadı")
            if not silent:
                messagebox.showerror(APP, f"Analiz hatası:\n{e}")
            return
        self.status.config(text="Hazır")
        info = summarize(self.templates)
        old = {f["name"]: f for f in self.fields}
        self.fields = [old[n] if n in old else new_field(n, d["auto"], "center" if self.mode == "pdf" else "left")
                       for n, d in info.items()]
        self.field_info = info
        stems = [t.stem for t in self.templates]
        if stems:
            self.lbl_tpl.config(text=f"{len(stems)} şablon: {', '.join(stems)}   •   {len(info)} alan bulundu")
        else:
            self.lbl_tpl.config(text=f"Bu klasörde {' / '.join(MODES[self.mode]['exts'])} şablon bulunamadı.")
        self.refresh_fields()
        self.refresh_table()
        if not silent and self.templates and not info:
            messagebox.showinfo(APP, "Şablonlarda değişken alan bulunamadı.\n\n" + MODE_HINT[self.mode])

    def analysis_progress(self, k, n, name):
        self.status.config(text=f"Şablonlar okunuyor ({k}/{n}): {name}"
                                + ("  –  Illustrator ile açılıyor, biraz sürebilir"
                                   if name.lower().endswith(AI_EXTS) else ""))
        self.update_idletasks()

    def _field_row(self, f):
        info = getattr(self, "field_info", {}).get(f["name"], {"files": [], "sample": "", "size": "", "tr": ""})
        stems = [t.stem for t in self.templates]
        miss = [s for s in stems if s not in info["files"]]
        return ("☑" if f["use"] else "☐", f["name"], ", ".join(info["files"]), ", ".join(miss) or "–",
                info["sample"], info["size"], info["tr"], ALIGN.get(f["align"], ""), CASES.get(f["case"], ""), f["max"] or "kapalı", f["minpct"],
                EMPTY.get(f["empty"], ""), f["prefix"], f["suffix"])

    def refresh_fields(self):
        sel = self.ftree.selection()
        self.ftree.delete(*self.ftree.get_children())
        for i, f in enumerate(self.fields):
            self.ftree.insert("", "end", iid=str(i), values=self._field_row(f))
        if sel and self.ftree.exists(sel[0]):
            self.ftree.selection_set(sel[0])
        elif self.fields:
            self.ftree.selection_set("0")
        else:
            self.lbl_field.config(text="(alan seçin)")
        self._pattern_buttons()

    def _field_click(self, ev):
        if self.ftree.identify_region(ev.x, ev.y) != "cell" or self.ftree.identify_column(ev.x) != "#1":
            return
        iid = self.ftree.identify_row(ev.y)
        if iid:
            f = self.fields[int(iid)]
            f["use"] = not f["use"]
            self.ftree.item(iid, values=self._field_row(f))
            self.refresh_table()
            self._pattern_buttons()

    def _field_select(self, ev=None):
        sel = self.ftree.selection()
        if not sel:
            return
        f = self.fields[int(sel[0])]
        self._loading = True
        self.lbl_field.config(text=f"Alan: {f['name']}")
        self.fv["case"].set(CASES.get(f["case"], CASES["asis"]))
        self.fv["empty"].set(EMPTY.get(f["empty"], EMPTY["blank"]))
        self.fv["align"].set(ALIGN.get(f["align"], ALIGN["left"]))
        self.fv["max"].set(str(f["max"]))
        self.fv["minpct"].set(str(f["minpct"]))
        self.fv["prefix"].set(f["prefix"])
        self.fv["suffix"].set(f["suffix"])
        self._loading = False

    def _field_changed(self, *a):
        sel = self.ftree.selection()
        if self._loading or not sel:
            return
        f = self.fields[int(sel[0])]
        f["case"] = next((k for k, v in CASES.items() if v == self.fv["case"].get()), "asis")
        f["empty"] = next((k for k, v in EMPTY.items() if v == self.fv["empty"].get()), "blank")
        f["align"] = next((k for k, v in ALIGN.items() if v == self.fv["align"].get()), "left")
        try:
            f["max"] = max(0, int(self.fv["max"].get()))
        except ValueError:
            pass
        try:
            f["minpct"] = min(100, max(1, int(float(self.fv["minpct"].get()))))
        except ValueError:
            pass
        f["prefix"] = self.fv["prefix"].get()
        f["suffix"] = self.fv["suffix"].get()
        self.ftree.item(sel[0], values=self._field_row(f))

    # ---- sekme 2: veri ----
    def _tab_data(self):
        tab = ttk.Frame(self.nb, padding=10)
        self.nb.add(tab, text=" 2. Veri ")
        bar = ttk.Frame(tab)
        bar.pack(fill="x")
        ttk.Button(bar, text="+ Satır ekle", command=self.add_row).pack(side="left")
        ttk.Button(bar, text="+ 10 satır", command=lambda: self.add_row(10)).pack(side="left", padx=(4, 0))
        ttk.Button(bar, text="Seçilenleri sil", command=self.delete_rows).pack(side="left", padx=4)
        ttk.Button(bar, text="Veriyi sıfırla", command=self.clear_rows).pack(side="left")
        ttk.Separator(bar, orient="vertical").pack(side="left", fill="y", padx=10)
        ttk.Button(bar, text="Toplu giriş…", style="Big.TButton", command=self.bulk_dialog).pack(side="left")
        ttk.Button(bar, text="Excel / CSV yükle…", command=self.import_file).pack(side="left", padx=4)
        ttk.Separator(bar, orient="vertical").pack(side="left", fill="y", padx=10)
        ttk.Button(bar, text="Boş şablon Excel indir…", command=lambda: self.export_xlsx(False)).pack(side="left")
        ttk.Button(bar, text="Tabloyu Excel'e aktar…", command=lambda: self.export_xlsx(True)).pack(side="left", padx=4)

        f = ttk.Frame(tab)
        f.pack(fill="both", expand=True, pady=(8, 0))
        self.tree = ttk.Treeview(f, show="headings", selectmode="extended")
        sy = ttk.Scrollbar(f, orient="vertical", command=self.tree.yview)
        sx = ttk.Scrollbar(f, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=sy.set, xscrollcommand=sx.set)
        sy.pack(side="right", fill="y")
        sx.pack(side="bottom", fill="x")
        self.tree.pack(fill="both", expand=True)
        self.tree.tag_configure("bad", foreground="#b00020")
        self.tree.bind("<Double-1>", self._cell_dblclick)
        self.tree.bind("<Delete>", lambda e: self.delete_rows())
        self.tree.bind("<Control-v>", lambda e: self.bulk_dialog(from_clipboard=True))
        self.tree.bind("<Control-a>", lambda e: self.tree.selection_set(self.tree.get_children()))
        self.lbl_rows = ttk.Label(tab, foreground="#555", text="")
        self.lbl_rows.pack(anchor="w", pady=(6, 0))

    def data_columns(self):
        return [COL_T] + [f["name"] for f in self.fields if f["use"]] + [COL_F]

    def refresh_table(self):
        self._close_editor()
        cols = ["#"] + self.data_columns()
        self.tree["columns"] = cols
        for c in cols:
            if c == "#":
                self.tree.heading(c, text=c)
            else:  # başlığa tıklamak o sütunun toplu girişini açar
                self.tree.heading(c, text=f"{c}   ▤ Toplu", command=lambda col=c: self.bulk_dialog(col))
            self.tree.column(c, width=50 if c == "#" else 170, anchor="center" if c == "#" else "w",
                             stretch=(c != "#"))
        self.tree.delete(*self.tree.get_children())
        lookup = build_lookup(self.templates)
        bad = 0
        for i, r in enumerate(self.rows):
            tv = str(r.get(COL_T, "") or "").strip()
            ok = (not self.templates) or (norm(tv) in lookup) or (not tv and len(self.templates) == 1)
            bad += 0 if ok else 1
            self.tree.insert("", "end", iid=str(i), values=[i + 1] + [r.get(c, "") for c in cols[1:]],
                             tags=() if ok else ("bad",))
        msg = (f"{len(self.rows)} satır.  Hücre için çift tıklayın (Tab: sonraki hücre). "
               f"Bir sütunu topluca doldurmak için başlığına tıklayın.")
        if bad:
            msg += f"   Kırmızı {bad} satırın şablonu klasörde yok."
        self.lbl_rows.config(text=msg)
        self.update_name_preview()

    def add_row(self, count=1):
        self.rows.extend({} for _ in range(count))
        self.refresh_table()
        self.tree.see(str(len(self.rows) - 1))
        iid = str(len(self.rows) - count)
        self.tree.see(iid)
        self.tree.selection_set(iid)
        self.after(50, lambda: self._start_edit(iid, 1))

    def delete_rows(self):
        idx = sorted((int(i) for i in self.tree.selection()), reverse=True)
        for i in idx:
            del self.rows[i]
        self.refresh_table()

    def clear_rows(self):
        if self.rows and messagebox.askyesno(APP, "Tablodaki tüm satırlar silinsin mi?"):
            self.rows = []
            self.refresh_table()

    def _cell_dblclick(self, ev):
        iid = self.tree.identify_row(ev.y)
        col = self.tree.identify_column(ev.x)
        if iid and col and col != "#1":
            self._start_edit(iid, int(col[1:]) - 1)

    def _close_editor(self):
        if self._editor is not None:
            ed, self._editor = self._editor, None
            ed.destroy()

    def _start_edit(self, iid, ci):
        self._close_editor()
        cols = self.tree["columns"]
        if not self.tree.exists(iid) or not (1 <= ci < len(cols)):
            return
        self.tree.see(iid)
        self.update_idletasks()
        bbox = self.tree.bbox(iid, column=cols[ci])
        if not bbox:
            return
        name = cols[ci]
        row = self.rows[int(iid)]
        var = tk.StringVar(value=str(row.get(name, "") or ""))
        if name == COL_T and self.templates:
            ed = ttk.Combobox(self.tree, textvariable=var, values=[t.stem for t in self.templates])
        else:
            ed = ttk.Entry(self.tree, textvariable=var)
        ed.place(x=bbox[0], y=bbox[1], width=bbox[2], height=bbox[3])
        ed.focus_set()
        ed.select_range(0, "end")
        self._editor = ed

        def commit(move=0):
            if self._editor is not ed:
                return
            row[name] = var.get().strip()
            self._editor = None
            ed.destroy()
            self.refresh_table()
            self.tree.selection_set(iid)
            if move:
                ni, nc = int(iid), ci + move
                if nc >= len(cols):
                    ni, nc = ni + 1, 1
                elif nc < 1:
                    ni, nc = ni - 1, len(cols) - 1
                if 0 <= ni < len(self.rows):
                    self.after(10, lambda: self._start_edit(str(ni), nc))
            return "break"

        ed.bind("<Return>", lambda e: commit())
        ed.bind("<Tab>", lambda e: commit(1))
        ed.bind("<Shift-Tab>", lambda e: commit(-1))
        ed.bind("<Escape>", lambda e: self._close_editor())
        ed.bind("<FocusOut>", lambda e: commit())

    def bulk_dialog(self, focus_col=None, from_clipboard=False):
        self._close_editor()
        text = ""
        if from_clipboard:
            try:
                text = self.clipboard_get()
            except tk.TclError:
                text = ""
        return BulkDialog(self, focus_col, text)

    def import_file(self):
        p = filedialog.askopenfilename(title="Veri dosyası seç", filetypes=[
            ("Excel / CSV / metin", "*.xlsx *.xlsm *.csv *.txt *.tsv"), ("Tümü", "*.*")])
        if not p:
            return
        try:
            rows = read_table_file(p)
        except ImportError:
            messagebox.showerror(APP, "Excel okumak için openpyxl gerekli.\n\nKomut satırında:  pip install openpyxl")
            return
        except Exception as e:
            messagebox.showerror(APP, f"Dosya okunamadı:\n{e}")
            return
        if not rows:
            messagebox.showinfo(APP, "Dosyada veri bulunamadı.")
            return
        ImportDialog(self, rows, f"İçe aktar – {Path(p).name}")

    def export_xlsx(self, with_rows):
        name = "veri.xlsx" if with_rows else "veri-sablonu.xlsx"
        p = filedialog.asksaveasfilename(title="Excel kaydet", defaultextension=".xlsx", initialfile=name,
                                         filetypes=[("Excel", "*.xlsx")])
        if not p:
            return
        try:
            write_xlsx(p, self.data_columns(), self.rows if with_rows else [],
                       [t.stem for t in self.templates], [f["name"] for f in self.fields if f["use"]])
        except ImportError:
            messagebox.showerror(APP, "Excel yazmak için openpyxl gerekli.\n\nKomut satırında:  pip install openpyxl")
            return
        except Exception as e:
            messagebox.showerror(APP, f"Kaydedilemedi (dosya Excel'de açık olabilir):\n{e}")
            return
        if messagebox.askyesno(APP, f"Kaydedildi:\n{p}\n\nŞimdi açılsın mı?"):
            open_path(p)

    # ---- sekme 3: isimlendirme ve seçenekler ----
    def _tab_options(self):
        tab = ttk.Frame(self.nb, padding=10)
        self.nb.add(tab, text=" 3. İsimlendirme ve Seçenekler ")
        g = ttk.LabelFrame(tab, text="Dosya isimlendirme", padding=8)
        g.pack(fill="x")
        g.columnconfigure(1, weight=1)
        ttk.Label(g, text="Dosya adı kalıbı:").grid(row=0, column=0, sticky="w")
        self.ent_pattern = ttk.Entry(g, textvariable=self.v_pattern, font=("Consolas", 11))
        self.ent_pattern.grid(row=0, column=1, sticky="ew", padx=6)
        self.lbl_ext = ttk.Label(g, text=".eps", width=5)
        self.lbl_ext.grid(row=0, column=2, sticky="w")
        ttk.Label(g, text="Alan ekle:").grid(row=1, column=0, sticky="w", pady=(6, 0))
        self.pbar = ttk.Frame(g)
        self.pbar.grid(row=1, column=1, columnspan=2, sticky="w", padx=6, pady=(6, 0))
        ttk.Label(g, text="Önizleme:").grid(row=2, column=0, sticky="w", pady=(6, 0))
        self.lbl_preview = ttk.Label(g, text="", foreground="#0b5394", font=("Consolas", 10))
        self.lbl_preview.grid(row=2, column=1, columnspan=2, sticky="w", padx=6, pady=(6, 0))
        ttk.Label(g, foreground="#555", text=(
            "{sıra:03} → 001, 002…   Kalıpta / kullanırsanız alt klasör oluşur (örn. {Beden}/{NAME}).")
                  ).grid(row=3, column=0, columnspan=3, sticky="w", pady=(6, 0))
        ttk.Checkbutton(g, text=f"Tablodaki '{COL_F}' sütunu doluysa kalıp yerine onu kullan",
                        variable=self.v_usename, command=self.update_name_preview).grid(
            row=4, column=0, columnspan=3, sticky="w", pady=(8, 0))
        ttk.Checkbutton(g, text="Her şablon için ayrı alt klasör aç (export/M/…, export/XL/…)",
                        variable=self.v_sub, command=self.update_name_preview).grid(
            row=5, column=0, columnspan=3, sticky="w")

        if self.mode == "eps":
            o = ttk.LabelFrame(tab, text="Çıktı biçimi", padding=8)
            o.pack(fill="x", pady=(10, 0))
            o.columnconfigure(5, weight=1)
            for j, (k, v) in enumerate(OUT_FORMATS.items()):
                ttk.Radiobutton(o, text=v, variable=self.v_fmt, value=k, command=self._fmt_changed).grid(
                    row=0, column=j, sticky="w", padx=(0, 18))
            ttk.Label(o, foreground="#555", text="EPS: fontlar gömülür.  PDF: Illustrator'da düzenlenebilir kalır.  "
                      "TIFF: Illustrator'dan dışa aktarılır (çalışma yüzeyi sınırında, LZW sıkıştırmalı)."
                      ).grid(row=0, column=4, columnspan=2, sticky="w", padx=(12, 0))
            tf = ttk.Frame(o)
            tf.grid(row=1, column=0, columnspan=6, sticky="w", pady=(8, 0))
            ttk.Label(tf, text="TIFF çözünürlüğü (dpi):").pack(side="left")
            self.sp_dpi = ttk.Spinbox(tf, textvariable=self.v_tif_dpi, from_=72, to=2400, increment=50, width=7)
            self.sp_dpi.pack(side="left", padx=(6, 18))
            ttk.Label(tf, text="Renk modu:").pack(side="left")
            self.cb_color = ttk.Combobox(tf, textvariable=self.v_tif_color, values=list(TIF_COLORS.values()),
                                         state="readonly", width=8)
            self.cb_color.pack(side="left", padx=6)
            self._fmt_changed()

        c = ttk.LabelFrame(tab, text="Çıktı klasöründe aynı isimde dosya varsa", padding=8)
        c.pack(fill="x", pady=(10, 0))
        for k, v in COLLISION.items():
            ttk.Radiobutton(c, text=v, variable=self.v_coll, value=k).pack(side="left", padx=(0, 18))

        s = ttk.LabelFrame(tab, text="Güvenlik ve rapor", padding=8)
        s.pack(fill="x", pady=(10, 0))
        ttk.Checkbutton(s, variable=self.v_protect, text=(
            "Orijinal şablon dosyalarını koru (şablon klasörüne ve şablonların üzerine yazmayı engelle)")
                        ).pack(anchor="w")
        ttk.Checkbutton(s, text="Üretim sonunda rapor dosyasını çıktı klasörüne kaydet",
                        variable=self.v_savereport).pack(anchor="w")
        ttk.Checkbutton(s, text="Üretim bitince çıktı klasörünü aç", variable=self.v_openafter).pack(anchor="w")
        self._pattern_buttons()

    def _fmt_changed(self):
        on = self.v_fmt.get() == "tif"
        self.sp_dpi.config(state="normal" if on else "disabled")
        self.cb_color.config(state="readonly" if on else "disabled")

    def tif_dpi(self):
        try:
            return min(2400, max(72, int(float(self.v_tif_dpi.get()))))
        except ValueError:
            return 300

    def tif_color(self):
        return next((k for k, v in TIF_COLORS.items() if v == self.v_tif_color.get()), "cmyk")

    def _pattern_buttons(self):
        if not hasattr(self, "pbar"):
            return
        for w in self.pbar.winfo_children():
            w.destroy()
        for ph in [COL_T] + [f["name"] for f in self.fields if f["use"]] + ["sıra:03", "tarih"]:
            ttk.Button(self.pbar, text="{" + ph + "}",
                       command=lambda p=ph: self.ent_pattern.insert("insert", "{" + p + "}")).pack(side="left", padx=2)

    def update_name_preview(self):
        if not hasattr(self, "lbl_preview"):
            return
        row = next((r for r in self.rows if any(str(v or "").strip() for v in r.values())), None)
        if row is None:
            row = {COL_T: self.templates[0].stem if self.templates else "M"}
            row.update({f["name"]: f["name"] for f in self.fields if f["use"]})
        t = build_lookup(self.templates).get(norm(row.get(COL_T, "")))
        tname = t.stem if t else str(row.get(COL_T, "") or "")
        custom = str(row.get(COL_F, "") or "").strip()
        if self.v_usename.get() and custom:
            parts, unknown = [clean_part(OUT_EXT_RE.sub("", custom))], []
        else:
            parts, unknown = name_parts(self.v_pattern.get(), row, 1, tname)
        if self.v_sub.get() and tname:
            parts = [clean_part(tname)] + parts
        ext = {"svg": ".svg", "pdf": ".pdf"}.get(self.mode, "." + self.v_fmt.get())
        self.lbl_ext.config(text=ext)
        txt = "\\".join(parts or ["satir_1"]) + ext
        if unknown:
            txt += "     ⚠ bilinmeyen alan: " + ", ".join("{" + u + "}" for u in unknown)
        self.lbl_preview.config(text=txt)

    # ---- sekme 4: üretim ----
    def _tab_run(self):
        tab = ttk.Frame(self.nb, padding=10)
        self.nb.add(tab, text=" 4. Üret ")
        bar = ttk.Frame(tab)
        bar.pack(fill="x")
        ttk.Button(bar, text="Kontrol Et (dosya yazmadan)", command=lambda: self.run(True)).pack(side="left")
        ttk.Button(bar, text="Seçili satırı önizle", command=self.preview_row).pack(side="left", padx=6)
        ttk.Button(bar, text="ÜRET", style="Big.TButton", command=lambda: self.run(False)).pack(side="left", padx=6)
        ttk.Button(bar, text="Çıktı klasörünü aç", command=lambda: self.open_dir(self.v_out)).pack(side="right")
        ttk.Button(bar, text="Raporu kaydet…", command=self.save_report).pack(side="right", padx=6)
        ttk.Checkbutton(tab, variable=self.v_checkdir, text=(
            f"Elle kontrol gerekenleri ayrı klasöre koymayı sor ({CHECK_DIR}): fontta olmayan harf "
            f"(ör. Türkçe karakter) veya karakter sınırını aşan yazı")).pack(anchor="w", pady=(8, 0))
        if self.mode == "eps":
            self.v_conv_dir = tk.StringVar()
            self.v_conv_rec = tk.BooleanVar(value=False)
            self.v_conv_skip = tk.BooleanVar(value=True)
            cv = ttk.LabelFrame(tab, text="Hazır EPS / AI dosyalarını TIFF'e çevir (Illustrator ile)", padding=8)
            cv.pack(fill="x", pady=(10, 0))
            cv.columnconfigure(1, weight=1)
            ttk.Label(cv, text="Klasör:").grid(row=0, column=0, sticky="w")
            ttk.Entry(cv, textvariable=self.v_conv_dir).grid(row=0, column=1, sticky="ew", padx=6)
            ttk.Button(cv, text="Gözat…", command=lambda: self.browse(self.v_conv_dir)).grid(row=0, column=2)
            ttk.Button(cv, text="Çıktı klasörü", command=lambda: self.v_conv_dir.set(self.v_out.get())).grid(
                row=0, column=3, padx=(4, 0))
            ttk.Button(cv, text="TIFF'e çevir", style="Big.TButton", command=self.convert_folder).grid(
                row=0, column=4, rowspan=2, padx=(10, 0))
            row = ttk.Frame(cv)
            row.grid(row=1, column=0, columnspan=4, sticky="w", pady=(6, 0))
            ttk.Checkbutton(row, text="Alt klasörler de", variable=self.v_conv_rec).pack(side="left")
            ttk.Checkbutton(row, text="TIFF'i zaten olanları atla", variable=self.v_conv_skip).pack(side="left", padx=12)
            ttk.Label(row, foreground="#555", text="Her dosyanın yanına aynı adlı .tif yazılır; çözünürlük ve renk "
                      "modu 3. sekmedeki TIFF ayarlarından alınır.").pack(side="left", padx=(6, 0))
        self.prog = ttk.Progressbar(tab, mode="determinate")
        self.prog.pack(fill="x", pady=8)
        f = ttk.Frame(tab)
        f.pack(fill="both", expand=True)
        self.log = tk.Text(f, wrap="word", font=("Consolas", 10), state="disabled")
        sb = ttk.Scrollbar(f, orient="vertical", command=self.log.yview)
        self.log.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.log.pack(fill="both", expand=True)

    def convert_folder(self):
        folder = self.v_conv_dir.get().strip() or self.v_out.get().strip()
        if not folder or not Path(folder).is_dir():
            messagebox.showwarning(APP, "Çevrilecek EPS'lerin bulunduğu klasörü seçin.")
            return
        self.v_conv_dir.set(folder)
        self.prog.config(value=0)

        def progress(k, n, name):
            self.prog.config(maximum=n, value=k)
            self.status.config(text=f"TIFF'e çevriliyor ({k}/{n}): {name}")
            self.update_idletasks()

        dpi, color = self.tif_dpi(), self.tif_color()
        rep = convert_tiff(folder, self.v_conv_rec.get(), dpi, color, self.v_conv_skip.get(), progress)
        L = [f"TIFF DÖNÜŞTÜRME – {datetime.datetime.now():%d.%m.%Y %H:%M}", f"Klasör: {folder}",
             f"Ayarlar: {dpi} dpi, {TIF_COLORS[color]}", "",
             f"Bulunan EPS/AI: {rep['total']}   Çevrilen: {len(rep['done'])}   "
             f"Atlanan (TIFF'i vardı): {len(rep['skipped'])}   Hata: {len(rep['errors'])}", ""]
        if rep["errors"]:
            L += ["HATALAR"] + ["  " + e for e in rep["errors"]] + [""]
        if rep["done"]:
            L += ["ÇEVRİLENLER"] + ["  " + Path(x).name for x in rep["done"]] + [""]
        if rep["skipped"]:
            L += ["ATLANANLAR"] + ["  " + Path(x).name for x in rep["skipped"]]
        self.last_report = "\n".join(L)
        self.set_log(self.last_report)
        self.status.config(text=f"TIFF dönüştürme tamam: {len(rep['done'])} dosya, {len(rep['errors'])} hata")
        if not rep["total"]:
            messagebox.showinfo(APP, "Bu klasörde EPS ya da AI dosyası bulunamadı.")

    def set_log(self, text):
        self.log.config(state="normal")
        self.log.delete("1.0", "end")
        self.log.insert("1.0", text)
        self.log.config(state="disabled")

    def options(self, out_dir=None):
        return {"out_dir": out_dir or self.v_out.get().strip(), "pattern": self.v_pattern.get(),
                "use_name_col": self.v_usename.get(), "subfolder": self.v_sub.get(),
                "collision": self.v_coll.get(), "protect": self.v_protect.get(), "fmt": self.v_fmt.get(),
                "tif_dpi": self.tif_dpi(), "tif_color": self.tif_color()}

    def preflight(self):
        self._close_editor()
        if not self.templates:
            self.run_analysis(silent=True)
        if not self.templates:
            messagebox.showwarning(APP, "Önce şablon klasörünü seçip analiz edin.")
            return False
        if not any(f["use"] for f in self.fields):
            messagebox.showwarning(APP, "Değişken olarak işaretlenmiş alan yok.")
            return False
        if not self.rows:
            messagebox.showwarning(APP, "Veri tablosu boş.")
            return False
        return True

    def run(self, dry):
        if not self.preflight():
            return
        out = self.v_out.get().strip()
        if not out:
            messagebox.showwarning(APP, "Çıktı klasörünü seçin.")
            return
        if self.v_protect.get() and Path(out).resolve() == Path(self.v_tpl.get().strip()).resolve():
            messagebox.showwarning(APP, "Çıktı klasörü şablon klasörüyle aynı.\n\n"
                                        "Orijinalleri korumak için farklı bir çıktı klasörü seçin "
                                        "veya 3. sekmeden korumayı kapatın.")
            return
        if not dry and not self.v_protect.get():
            if not messagebox.askyesno(APP, "Orijinal dosya koruması KAPALI. Şablon dosyalarının üzerine "
                                            "yazılabilir. Devam edilsin mi?", icon="warning"):
                return
        self.nb.select(3)
        self.prog.config(value=0)

        def progress(k, n):
            self.prog.config(maximum=n, value=k)
            self.update_idletasks()

        # Şablonlar diskte değişmiş olabilir; üretimden önce yeniden oku.
        try:
            self.templates = analyze(self.v_tpl.get().strip(), self.v_rec.get(), out,
                                     progress=self.analysis_progress, mode=self.mode)
        except Exception as e:
            messagebox.showerror(APP, f"Şablonlar okunamadı:\n{e}")
            return
        opts = dict(self.options(), check_dir=False)
        if self.v_checkdir.get():
            # Önce yazmadan dene: kontrol gereken dosya varsa nereye konacağını sor.
            flagged = generate(opts, self.templates, self.fields, self.rows, dry=True)["check"]
            if flagged and dry:
                opts["check_dir"] = True
            elif flagged:
                names = ", ".join(label for _, label, _, _ in flagged[:8]) + ("…" if len(flagged) > 8 else "")
                ans = messagebox.askyesnocancel(APP, (
                    f"{len(flagged)} dosyanın elle kontrol edilmesi gerekiyor (fontta olmayan harf veya uzun "
                    f"yazı):\n\n{names}\n\nBunlar '{CHECK_DIR}' adlı ayrı klasöre konsun mu?\n\n"
                    f"Evet: ayrı klasöre.   Hayır: diğerleriyle aynı yere.   İptal: üretme."))
                if ans is None:
                    return
                opts["check_dir"] = ans
        rep = generate(opts, self.templates, self.fields, self.rows, dry=dry, progress=progress)
        self.last_report = report_text(rep, out)
        self.set_log(self.last_report)
        self.status.config(text=("Kontrol tamam: " if dry else "Üretim tamam: ") +
                                f"{len(rep['done'])} dosya, {len(rep['missing'])} eksik şablon, "
                                f"{len(rep['check'])} kontrol gerekli")
        if dry:
            return
        if self.v_savereport.get() and rep["done"]:
            try:
                rp = Path(out) / f"rapor_{datetime.datetime.now():%Y%m%d_%H%M%S}.txt"
                rp.write_text(self.last_report, encoding="utf-8-sig")
            except OSError:
                pass
        if self.v_openafter.get() and Path(out).is_dir():
            open_path(out)

    def preview_row(self):
        if not self.preflight():
            return
        sel = self.tree.selection()
        row = self.rows[int(sel[0])] if sel else self.rows[0]
        tmp = Path(tempfile.gettempdir()) / "grafik-varyant-onizleme" / self.mode
        opts = dict(self.options(str(tmp)), collision="overwrite", protect=False, subfolder=False)
        rep = generate(opts, self.templates, self.fields, [row])
        if rep["done"]:
            open_path(rep["done"][0][3])
        else:
            messagebox.showwarning(APP, "Önizleme üretilemedi:\n\n" + report_text(rep))

    def save_report(self):
        if not self.last_report:
            messagebox.showinfo(APP, "Henüz rapor yok.")
            return
        p = filedialog.asksaveasfilename(title="Raporu kaydet", defaultextension=".txt", initialfile="rapor.txt",
                                         filetypes=[("Metin", "*.txt")])
        if p:
            Path(p).write_text(self.last_report, encoding="utf-8-sig")


def main():
    try:
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass
    App().mainloop()


if __name__ == "__main__":
    main()
